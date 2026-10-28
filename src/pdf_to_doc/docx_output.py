"""Renders the document model as a Word file for editing, keeping its structure rather than its page
layout. A page kept as an image is the exception: a Word page of its own, the image filling it."""

import io
import re
from collections.abc import Iterator
from copy import deepcopy

import docx
from docx.document import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.section import CT_SectPr
from docx.oxml.text.paragraph import CT_P
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Cm, Emu, Length, Mm, Pt
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph

from pdf_to_doc.model import (
    DEEPEST_HEADING,
    LARGEST_PAGE,
    Block,
    Document,
    Heading,
    KeptImage,
    ListItem,
    Placeholder,
    Span,
    Table,
    TableCell,
    stretches,
    uncertain_text,
)
from pdf_to_doc.summary import Summary
from pdf_to_doc.thai import THAI_LETTER

# Thai and English alike, so lines that mix them keep one size.
_FONT = "TH Sarabun New"
_BODY_SIZE = 16  # points
# Heading sizes by level; deeper levels are bold at the body size.
_HEADING_SIZES = {1: 20, 2: 18}
# The space the template leaves after each paragraph, kept below a table too.
_SPACE_AFTER_TABLE = Pt(10)
_LARGEST_PAGE = Pt(LARGEST_PAGE)  # Word takes no page larger either way
# A4 with Word's usual margins, in place of the template's US Letter, and the text area between them.
_A4_WIDTH, _A4_HEIGHT = Mm(210), Mm(297)
_MARGIN = Cm(2.54)
_TEXT_WIDTH, _TEXT_HEIGHT = Emu(_A4_WIDTH - 2 * _MARGIN), Emu(_A4_HEIGHT - 2 * _MARGIN)

# A run of English: from a letter or digit to a letter or digit, holding ASCII spaces
# and punctuation, typographic apostrophes and the en dash in between.
_LATIN = re.compile(r"[A-Za-z0-9](?:[ -~\u2018\u2019\u2013]*[A-Za-z0-9])?")
_LETTER = re.compile(r"[A-Za-z]")
_THAI = re.compile(THAI_LETTER)


def render_docx(document: Document, summary: Summary) -> bytes:
    word = docx.Document()
    [section] = word.sections
    section.page_width, section.page_height = _A4_WIDTH, _A4_HEIGHT
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = _MARGIN
    _set_default_font(word.styles.element)
    for level in range(1, DEEPEST_HEADING + 1):
        _set_heading_font(word.styles[f"Heading {level}"].element, _HEADING_SIZES.get(level, _BODY_SIZE))
    unmarked_list_id = _add_unmarked_list(word.part.numbering_part.element)
    a4 = deepcopy(section._sectPr)  # the section properties of every page not kept as an image
    summary_paragraph = word.add_paragraph()
    _add_text(summary_paragraph, "สรุปการแปลง", bold=True)
    _add_text(summary_paragraph, f" {summary.headline(uncertain_marked_as='ไฮไลต์สีเหลือง')}")
    writing = a4  # the properties of the section being written
    for index, page in enumerate(document.pages):
        kept = page.kept_image
        # A kept page is a section of its own, and so are the other pages after it, back on A4.
        # A section starts a new page, so its first paragraph needs no page break.
        new_section = kept is not None or writing is not a4
        if new_section:
            _end_section(word, writing)
        writing = a4 if kept is None else _kept_page_section(a4, kept)
        if kept is not None:
            _add_page_image(word, kept, writing)
            continue
        first: WordParagraph | None = None
        previous: Block | None = None
        for block in page.blocks:
            if isinstance(previous, Table) and isinstance(block, Table):
                word.add_paragraph()  # Word joins tables that touch into one
            paragraph = _add_block(word, block, unmarked_list_id)
            first = first or paragraph
            if isinstance(previous, Table) and not isinstance(block, (Heading, Table)):
                # Headings have space above them already; other text would sit on the table's bottom line.
                paragraph.paragraph_format.space_before = _SPACE_AFTER_TABLE
            previous = block
        if new_section and first is None:
            word.add_paragraph()  # an empty page, so that the section holds something
        elif index > 0 and not new_section:
            # Each page of the original starts a new page, at its own first paragraph. A page
            # break in a paragraph of its own falls onto a page of its own when the page before
            # is full, and leaves that page blank. An empty page is an empty paragraph.
            (first or word.add_paragraph()).paragraph_format.page_break_before = True
    if writing is not a4:
        # The last section's properties are the body's own.
        section._sectPr.getparent().replace(section._sectPr, writing)
    # Written to memory: a conversion leaves no file behind.
    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def _end_section(word: WordDocument, properties: CT_SectPr) -> None:
    """End the section written so far, as Word does: its properties go in its last paragraph. A
    section that ends with a table gets an empty paragraph after it, to carry them."""
    body = word.element.body
    last = body[-2]  # after it come the properties of the body's last section
    if not isinstance(last, CT_P):
        last = word.add_paragraph()._p
    last.set_sectPr(deepcopy(properties))


def _kept_page_section(a4: CT_SectPr, image: KeptImage) -> CT_SectPr:
    """The section properties of a kept page, from those of the other pages: its own size and
    orientation, with no margins. A page larger than Word takes is scaled down to fit, keeping its shape."""
    properties = deepcopy(a4)
    properties.page_width, properties.page_height = _shrunk(image, _LARGEST_PAGE, _LARGEST_PAGE)
    properties.orientation = WD_ORIENT.LANDSCAPE if image.width > image.height else WD_ORIENT.PORTRAIT
    for margin in ("top_margin", "right_margin", "bottom_margin", "left_margin", "header", "footer", "gutter"):
        setattr(properties, margin, 0)
    return properties


def _add_page_image(word: WordDocument, image: KeptImage, properties: CT_SectPr) -> None:
    """The image of a kept page, filling the page its section `properties` give. Its paragraph has
    no space around it, so that nothing is left over for a page of its own."""
    paragraph = word.add_paragraph()
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    # As large as the page itself, which Word sizes in whole twentieths of a point.
    paragraph.add_run().add_picture(io.BytesIO(image.data), width=properties.page_width, height=properties.page_height)


def _add_block(word: WordDocument, block: Block, unmarked_list_id: int) -> WordParagraph:
    """Adds the block, and returns the paragraph it starts with: for a table, its first cell's."""
    if isinstance(block, Table):
        first_paragraph: WordParagraph = _add_table(word, block).cell(0, 0).paragraphs[0]
        return first_paragraph
    if isinstance(block, Placeholder) and block.image is not None:
        return _add_picture(word, block.image)
    if isinstance(block, Heading):
        paragraph = word.add_paragraph(style=f"Heading {block.level}")
    elif isinstance(block, ListItem) and block.bulleted:
        paragraph = word.add_paragraph(style="List Bullet")
    elif isinstance(block, ListItem):
        paragraph = word.add_paragraph(style="List Paragraph")
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val = 0
        numbering.get_or_add_numId().val = unmarked_list_id
    else:
        paragraph = word.add_paragraph()
    _add_text(paragraph, block.text, uncertain=uncertain_text(block))
    return paragraph


def _add_picture(word: WordDocument, image: KeptImage) -> WordParagraph:
    """A picture kept as itself, in a paragraph of its own, at the size it is shown on the page;
    shrunk to the text area, keeping its shape, when it is larger. Single spacing keeps the line
    no taller than the picture."""
    width, height = _shrunk(image, _TEXT_WIDTH, _TEXT_HEIGHT)
    paragraph = word.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run().add_picture(io.BytesIO(image.data), width=width, height=height)
    return paragraph


def _shrunk(image: KeptImage, most_width: Length, most_height: Length) -> tuple[Length, Length]:
    """The size the image is shown at, shrunk, keeping its shape, when it is wider or taller than the most."""
    scale = min(1.0, most_width / Pt(image.width), most_height / Pt(image.height))
    return Pt(image.width * scale), Pt(image.height * scale)


def _add_table(word: WordDocument, table: Table) -> WordTable:
    """A Word table with real merged cells, its columns of equal width."""
    placed = list(_place_cells(table))
    columns = max(column + cell.colspan for _, column, cell in placed)
    word_table: WordTable = word.add_table(rows=len(table.rows), cols=columns)
    word_table.style = "Table Grid"
    for row, column, cell in placed:
        target = word_table.cell(row, column)
        if cell.rowspan > 1 or cell.colspan > 1:
            last_row = min(row + cell.rowspan, len(table.rows)) - 1
            target = target.merge(word_table.cell(last_row, column + cell.colspan - 1))
        _add_text(target.paragraphs[0], cell.text, uncertain=uncertain_text(cell))
    return word_table


def _place_cells(table: Table) -> Iterator[tuple[int, int, TableCell]]:
    """Each cell with the row and column it starts at. A cell starts at the first column
    of its row that no merged cell from the rows above reaches into."""
    covered: set[tuple[int, int]] = set()
    for row, cells in enumerate(table.rows):
        column = 0
        for cell in cells:
            while (row, column) in covered:
                column += 1
            yield row, column, cell
            covered.update(
                (row + down, column + across) for down in range(cell.rowspan) for across in range(cell.colspan)
            )
            column += cell.colspan


def _add_unmarked_list(numbering: BaseOxmlElement) -> int:
    """A list whose items show nothing of Word's before them, for items that keep their own
    number or letter as text. Their text lines up with that of bulleted items. Returns its numbering id."""
    taken = [int(abstract.get(qn("w:abstractNumId"))) for abstract in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(taken) + 1
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:multiLevelType w:val="singleLevel"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="none"/>'
        # No tab after the empty number, or the text would start at the next tab stop.
        '<w:suff w:val="nothing"/><w:lvlText w:val=""/><w:lvlJc w:val="left"/>'
        '<w:pPr><w:ind w:left="360" w:hanging="0"/></w:pPr></w:lvl>'
        "</w:abstractNum>"
    )
    # Every abstract definition comes before the first numbering instance.
    numbering.find(qn("w:num")).addprevious(abstract)
    return int(numbering.add_num(abstract_id).numId)


def _set_default_font(styles: BaseOxmlElement) -> None:
    """Every style starts from TH Sarabun New in all four font slots, Thai as the complex-script language."""
    run_properties = styles.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    _set_font(run_properties, _BODY_SIZE)
    language = run_properties.find(qn("w:lang"))
    language.set(qn("w:bidi"), "th-TH")


def _set_heading_font(style: BaseOxmlElement, size: float) -> None:
    """Bold TH Sarabun New in both slots, black and upright, in place of the template's theme fonts and colours."""
    run_properties = style.get_or_add_rPr()
    for child in list(run_properties):
        run_properties.remove(child)
    _set_font(run_properties, size)
    for index, tag in enumerate(("w:b", "w:bCs"), start=1):
        run_properties.insert(index, OxmlElement(tag))


def _set_font(run_properties: BaseOxmlElement, size: float) -> None:
    """Names the font in every slot, so no theme font applies, and sets the plain and complex-script sizes."""
    fonts = run_properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.attrib.clear()
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(slot), _FONT)
    for tag in ("w:sz", "w:szCs"):
        element = run_properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            run_properties.append(element)
        element.set(qn("w:val"), str(round(size * 2)))  # half-points


def _add_text(paragraph: WordParagraph, text: str, *, bold: bool = False, uncertain: tuple[Span, ...] = ()) -> None:
    """Thai goes in runs marked complex script, so Word breaks its lines between words. English
    words get runs of their own without the mark, or Word would proof them as Thai. Uncertain
    text is highlighted yellow."""
    for highlighted, stretch in stretches(text, uncertain):
        for latin, piece in _pieces(stretch):
            run = paragraph.add_run(piece)
            if not latin and _THAI.search(piece):
                run.font.complex_script = True
            if bold:
                # Thai is drawn with the complex-script bold, other text with the plain one.
                run.font.bold = run.font.cs_bold = True
            if highlighted:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _pieces(text: str) -> list[tuple[bool, str]]:
    """The text cut into English runs and the rest, in order. Numbers alone stay with the text around them."""
    pieces: list[tuple[bool, str]] = []
    position = 0
    for match in _LATIN.finditer(text):
        if not _LETTER.search(match[0]):
            continue
        if match.start() > position:
            pieces.append((False, text[position : match.start()]))
        pieces.append((True, match[0]))
        position = match.end()
    if position < len(text):
        pieces.append((False, text[position:]))
    return pieces
