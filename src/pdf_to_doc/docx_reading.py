"""Reads a DOCX source straight into the document model, from the structure Word saved.

Pages are where Word's last layout broke them, which Word records on saving. A file
without those records, as other programs write, breaks pages where breaks were typed
or set on a paragraph. Headers and footers are page furniture and are not read. Text
the author deleted with tracked changes on is left out, as is hidden text, and field
codes give way to what they show.
"""

import io
import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, replace
from enum import Enum, auto
from typing import cast

import docx
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from lxml import etree

from pdf_to_doc.docx_numbering import Numbering
from pdf_to_doc.intake import MAX_PAGES, DamagedFile, TooLarge
from pdf_to_doc.model import (
    DEEPEST_HEADING,
    Block,
    Document,
    Heading,
    ListItem,
    Page,
    Paragraph,
    Placeholder,
    Table,
    TableCell,
)
from pdf_to_doc.thai import repair

_MARKUP_COMPATIBILITY = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_ALTERNATE_CONTENT = f"{{{_MARKUP_COMPATIBILITY}}}AlternateContent"
_CHOICE = f"{{{_MARKUP_COMPATIBILITY}}}Choice"
_VML_IMAGE = "{urn:schemas-microsoft-com:vml}imagedata"
# Parts of a paragraph that hold no text a reader sees: its properties, text deleted
# with tracked changes on, and the codes behind fields.
_HIDDEN = frozenset(qn(tag) for tag in ("w:pPr", "w:rPr", "w:del", "w:moveFrom", "w:instrText", "w:delText"))
# Typed tabs and line breaks set text apart, as a space does.
_SPACING = frozenset(qn(tag) for tag in ("w:tab", "w:ptab", "w:br", "w:cr"))
# What a drawing shows when it is a picture: a bitmap, a chart or a diagram.
_PICTURES = (qn("a:blip"), _VML_IMAGE, qn("c:chart"), qn("dgm:relIds"))
# Elements that only wrap the paragraphs, tables, rows or cells inside them.
_WRAPPERS = frozenset(qn(tag) for tag in ("w:sdt", "w:sdtContent", "w:customXml"))
_SPACES = re.compile(r"[ \t]+")


class _Mark(Enum):
    PICTURE = auto()
    NEW_PAGE = auto()  # a page Word's layout began, or a typed page break
    BREAK_BEFORE = auto()  # a paragraph set to start a page: no new one when a page has just begun
    TEXT = auto()  # in a table cell: text was here, so a later mark comes after the table's start


# What reading a paragraph or table gives: blocks, and marks among them.
_Item = Block | _Mark


@dataclass(frozen=True)
class _TextBox:
    content: BaseOxmlElement  # its w:txbxContent, which holds paragraphs and tables


def read_docx(source: bytes) -> Document:
    """The DOCX's document model. A file that is not a whole DOCX raises `DamagedFile`, and one
    of more pages than the limit, as Word last laid it out, raises `TooLarge`."""
    try:
        word = docx.Document(io.BytesIO(source))
    except (KeyError, ValueError, zipfile.BadZipFile, etree.XMLSyntaxError, PackageNotFoundError) as error:
        raise DamagedFile() from error
    # Word records every page its layout began, typed breaks included; those alone count when they are there.
    laid_out = next(word.element.body.iter(qn("w:lastRenderedPageBreak")), None) is not None
    reader = _Reader(word, laid_out)
    pages = _paginate([item for element in _block_elements(word.element.body) for item in reader.read(element)])
    if len(pages) > MAX_PAGES:
        raise TooLarge("pages", len(pages), MAX_PAGES)
    return Document(pages=tuple(Page(number, "text", tuple(blocks)) for number, blocks in enumerate(pages, 1)))


def _paginate(items: list[_Item]) -> list[list[Block]]:
    """The blocks of each page. Every new-page mark starts a page, even one that only a
    table runs onto or that two typed breaks leave empty, as Word counts them; a paragraph
    set to start a page starts one unless a page has just begun."""
    pages: list[list[Block]] = [[]]
    for item in items:
        if item is _Mark.NEW_PAGE:
            pages.append([])
        elif item is _Mark.BREAK_BEFORE:
            if pages[-1]:
                pages.append([])
        elif not isinstance(item, _Mark):
            pages[-1].append(item)
    return pages


class _Reader:
    def __init__(self, word: WordDocument, laid_out: bool) -> None:
        self._laid_out = laid_out
        self._styles = _Styles(word.styles.element)
        try:
            definitions = word.part.numbering_part.element
        except NotImplementedError:  # python-docx cannot make the numbering part that a document without lists lacks
            definitions = None
        self._numbering = Numbering(definitions, self._styles.list_style_instance)

    def read(self, element: BaseOxmlElement) -> list[_Item]:
        """The blocks a paragraph or table becomes, with marks where new pages start.

        A picture in a table's cell is placed right after the table. A table that Word
        began on a new page starts that page; the pages it runs onto follow it.
        """
        if element.tag != qn("w:tbl"):
            return self._paragraph(element)
        held: list[_Item] = []
        table = self._table(element, held)
        placeholders = [item for item in held if isinstance(item, Placeholder)]
        pages = [item for item in held if item is _Mark.NEW_PAGE or item is _Mark.TEXT]
        before = 1 if pages and pages[0] is _Mark.NEW_PAGE else 0
        after = pages.count(_Mark.NEW_PAGE) - before
        return [*[_Mark.NEW_PAGE] * before, table, *placeholders, *[_Mark.NEW_PAGE] * after]

    def _paragraph(self, paragraph: BaseOxmlElement) -> list[_Item]:
        """The paragraph as a heading, a list item or a paragraph, cut where a picture, a text
        box or a new page sits.

        A list item keeps its number as text, as on a text page; a bullet is dropped.
        Only the first piece is a heading or list item; the rest are paragraphs.
        """
        listed = self._styles.numbering(paragraph)
        # An empty list paragraph still counts in its list, as Word counts it.
        label = self._numbering.label(*listed) if listed else None
        level = self._styles.outline_level(paragraph)
        items: list[_Item] = []
        if not self._laid_out and self._styles.breaks_page_before(paragraph):
            items.append(_Mark.BREAK_BEFORE)
        pieces: list[str] = []
        first = True

        def flush() -> None:
            nonlocal first
            text = _clean("".join(pieces))
            pieces.clear()
            if not text:
                return
            if first and label and label.text and not label.bullet:
                text = f"{label.text}{' ' if label.spaced else ''}{text}"
            if first and level is not None:
                items.append(Heading(min(level + 1, DEEPEST_HEADING), text))
            elif first and label and listed:
                items.append(ListItem(text, bulleted=label.bullet, level=listed[1]))
            else:
                items.append(Paragraph(text))
            first = False

        for piece in _pieces(paragraph, self._laid_out):
            if isinstance(piece, str):
                pieces.append(piece)
                continue
            flush()
            if isinstance(piece, _TextBox):
                # Where Word laid out a text box starts no page of the document's own.
                content = [item for element in _block_elements(piece.content) for item in self.read(element)]
                items += [item for item in content if not isinstance(item, _Mark)]
            elif piece is _Mark.PICTURE:
                items.append(Placeholder("image"))
            else:
                items.append(piece)
        flush()
        return items

    def _table(self, table: BaseOxmlElement, held: list[_Item]) -> Table:
        """The table's cells with their spans; each row holds the cells that start in it.
        `held` collects, in order, the placeholders and page marks read from the cells, with a
        text mark wherever there was text."""
        rows: list[list[TableCell]] = []
        # For each grid column, where the cell starts that a vertical merge carries on.
        starts: dict[int, tuple[int, int]] = {}
        for row in _children(table, "w:tr"):
            column = _setting(row, "w:trPr/w:gridBefore")
            starting = [TableCell("") for _ in range(column)]
            for cell in _children(row, "w:tc"):
                span = max(1, _setting(cell, "w:tcPr/w:gridSpan"))
                merge = cell.find(_path("w:tcPr/w:vMerge"))
                if merge is not None and merge.get(qn("w:val"), "continue") == "continue" and column in starts:
                    above_row, above_index = starts[column]
                    above = rows[above_row][above_index]
                    rows[above_row][above_index] = replace(above, rowspan=above.rowspan + 1)
                else:
                    starts[column] = (len(rows), len(starting))
                    starting.append(TableCell(self._cell(cell, held), colspan=span))
                column += span
            rows.append(starting)
        return Table(rows=tuple(tuple(row) for row in rows))

    def _cell(self, cell: BaseOxmlElement, held: list[_Item]) -> str:
        """A cell's text: its paragraphs, and those of any table inside it, joined with a space.
        What it holds besides text goes into `held`."""
        texts: list[str] = []
        for element in _block_elements(cell):
            for item in self.read(element):
                if isinstance(item, Table):
                    texts += [inner.text for row in item.rows for inner in row if inner.text]
                    held.append(_Mark.TEXT)
                elif isinstance(item, (Paragraph, Heading, ListItem)):
                    texts.append(item.text)
                    held.append(_Mark.TEXT)
                else:
                    held.append(item)
        return " ".join(texts)


def _block_elements(container: BaseOxmlElement) -> Iterator[BaseOxmlElement]:
    """The paragraphs and tables in a body, cell or text box, in order, inside content controls too."""
    for child in container:
        if child.tag in (qn("w:p"), qn("w:tbl")):
            yield child
        elif child.tag in _WRAPPERS:
            yield from _block_elements(child)


def _children(element: BaseOxmlElement, tag: str) -> Iterator[BaseOxmlElement]:
    """The element's children of one kind, inside content controls too, as Word's repeating sections hold rows."""
    for child in element:
        if child.tag == qn(tag):
            yield child
        elif child.tag in _WRAPPERS:
            yield from _children(child, tag)


def _clean(text: str) -> str:
    """The text repaired, its runs of spaces and tabs made single spaces."""
    return repair(_SPACES.sub(" ", text).strip())


def _pieces(element: BaseOxmlElement, laid_out: bool) -> Iterator[str | _Mark | _TextBox]:
    """The paragraph's text, piece by piece, with marks where pictures, text boxes and new pages sit.

    In a file Word laid out, new pages start where it recorded them; otherwise at typed page breaks.
    """
    for child in element:
        if child.tag == qn("w:t"):
            yield child.text or ""
        elif child.tag == qn("w:lastRenderedPageBreak"):
            yield _Mark.NEW_PAGE
        elif child.tag == qn("w:br") and child.get(qn("w:type")) == "page":
            if not laid_out:
                yield _Mark.NEW_PAGE
        elif child.tag in _SPACING:
            yield " "
        elif child.tag == qn("w:noBreakHyphen"):
            yield "-"
        elif child.tag in (qn("w:drawing"), qn("w:pict")):
            yield from _drawing(child)
        elif child.tag == _ALTERNATE_CONTENT:
            # Word writes a drawing twice: as itself, and as a fallback for older readers.
            choice = child.find(_CHOICE)
            if choice is not None:
                yield from _pieces(choice, laid_out)
        elif child.tag == qn("w:r") and _hidden(child):
            continue
        elif child.tag not in _HIDDEN:
            yield from _pieces(child, laid_out)  # runs, hyperlinks, insertions, field results and the like


def _hidden(run: BaseOxmlElement) -> bool:
    """Whether the run is formatted as hidden, which Word neither shows nor prints."""
    vanish = run.find(_path("w:rPr/w:vanish"))
    return vanish is not None and vanish.get(qn("w:val"), "true") not in ("0", "false", "off")


def _drawing(drawing: BaseOxmlElement) -> Iterator[_Mark | _TextBox]:
    """The text of each text box, or a mark for a picture. A drawn line or shape alone shows nothing to transcribe."""
    boxes = list(drawing.iter(qn("w:txbxContent")))
    if boxes:
        yield from (_TextBox(content) for content in boxes)
    elif any(next(drawing.iter(tag), None) is not None for tag in _PICTURES):
        yield _Mark.PICTURE


def _path(path: str) -> str:
    """An element path such as "w:pPr/w:numPr" in the namespaced form lxml finds."""
    return "/".join(qn(part) for part in path.split("/"))


def _setting(element: BaseOxmlElement, path: str) -> int:
    setting = element.find(_path(path))
    return 0 if setting is None else int(setting.get(qn("w:val"), "0"))


class _Styles:
    """The paragraph styles, looked up the way Word does: a paragraph's own setting
    first, then its style, then the styles that style is based on."""

    def __init__(self, styles: BaseOxmlElement) -> None:
        self._by_id = {style.get(qn("w:styleId")): style for style in styles.findall(qn("w:style"))}
        self._default = next(
            (
                style_id
                for style_id, style in self._by_id.items()
                if style.get(qn("w:type")) == "paragraph" and style.get(qn("w:default")) in ("1", "true")
            ),
            None,
        )

    def numbering(self, paragraph: BaseOxmlElement) -> tuple[str, int, str | None] | None:
        """The list instance the paragraph is numbered in, its level, and its style when the style
        rather than the paragraph puts it in the list; None when it is in no list."""
        own = paragraph.find(_path("w:pPr/w:numPr/w:numId"))
        instance = self._paragraph_setting(paragraph, "w:numPr/w:numId")
        if instance is None or instance.get(qn("w:val")) in (None, "0"):  # instance 0 takes a paragraph out of its list
            return None
        level = self._paragraph_setting(paragraph, "w:numPr/w:ilvl")
        ilvl = int(level.get(qn("w:val"), "0")) if level is not None else 0
        return str(instance.get(qn("w:val"))), ilvl, None if own is not None else self._style_of(paragraph)

    def list_style_instance(self, style_id: str) -> str | None:
        """The numbering instance a list style names."""
        style = self._by_id.get(style_id)
        instance = None if style is None else style.find(_path("w:pPr/w:numPr/w:numId"))
        return None if instance is None else str(instance.get(qn("w:val")))

    def breaks_page_before(self, paragraph: BaseOxmlElement) -> bool:
        setting = self._paragraph_setting(paragraph, "w:pageBreakBefore")
        return setting is not None and setting.get(qn("w:val"), "true") not in ("0", "false", "off")

    def outline_level(self, paragraph: BaseOxmlElement) -> int | None:
        """The paragraph's heading level, counting from 0, or None for body text."""
        setting = self._paragraph_setting(paragraph, "w:outlineLvl")
        if setting is None:
            return None
        level = int(setting.get(qn("w:val"), "9"))
        return level if level < 9 else None  # level 9 is body text

    def _style_of(self, paragraph: BaseOxmlElement) -> str | None:
        reference = paragraph.find(_path("w:pPr/w:pStyle"))
        return reference.get(qn("w:val")) if reference is not None else self._default

    def _paragraph_setting(self, paragraph: BaseOxmlElement, path: str) -> BaseOxmlElement | None:
        """The paragraph's own setting at `path` in its properties, or else its style chain's."""
        own = paragraph.find(_path(f"w:pPr/{path}"))
        if own is not None:
            return cast(BaseOxmlElement, own)
        style_id = self._style_of(paragraph)
        seen: set[str] = set()  # a style based on itself, however far round, ends the search
        while style_id and style_id not in seen and (style := self._by_id.get(style_id)) is not None:
            seen.add(style_id)
            setting = style.find(_path(f"w:pPr/{path}"))
            if setting is not None:
                return cast(BaseOxmlElement, setting)
            based_on = style.find(qn("w:basedOn"))
            style_id = based_on.get(qn("w:val")) if based_on is not None else None
        return None
