"""Builds small Word files from scratch, so tests need no Word install.

- `docx` writes a minimal DOCX: the three parts Word needs to open a one-paragraph document.
- `document` writes a DOCX with Word's own structure: headings in heading styles,
  lists whose numbers Word draws, tables with merged cells, pictures, and page
  breaks, typed or where Word's last layout broke the page.
- `legacy_doc` writes a Word 97-2003 .doc container: a compound file holding a
  WordDocument stream that starts with the Word 97 file information block. The
  stream holds nothing else, so the file is real enough to identify but Word
  would not open it.
- `locked_docx` writes the compound file Word saves a password-locked DOCX in,
  with its two streams named and empty.
- `docx_unpacking_to` writes a minimal DOCX padded to unpack to a given size, as a zip bomb.

The same inputs always give the same bytes.
"""

import io
import struct
import zipfile
from dataclasses import dataclass
from typing import Literal
from xml.sax.saxutils import escape

import docx as python_docx
import pymupdf
from docx.document import Document as WordDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from pdf_builder import Cell

_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml"'
    ' ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_RELATIONSHIPS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1"'
    ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"'
    ' Target="word/document.xml"/>'
    "</Relationships>"
)


def docx(text: str) -> bytes:
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p></w:body>"
        "</w:document>"
    )
    parts = {"[Content_Types].xml": _CONTENT_TYPES, "_rels/.rels": _RELATIONSHIPS, "word/document.xml": document}
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in parts.items():
            archive.writestr(_entry(name), content, zipfile.ZIP_DEFLATED)
    return buffer.getvalue()


def _entry(name: str) -> zipfile.ZipInfo:
    """A zip entry that is the same on every run and every system: Python otherwise dates
    an entry now, and records the system that made it, Windows or Unix."""
    entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
    entry.create_system = 0
    return entry


# Compound file layout (version 3, 512-byte sectors): the header, then sector 0
# holds the allocation table, sector 1 the directory, and sectors 2-9 the stream.
_SECTOR = 512
_FREE = 0xFFFFFFFF
_END_OF_CHAIN = 0xFFFFFFFE
_FAT_SECTOR = 0xFFFFFFFD
_NO_STREAM = 0xFFFFFFFF
# Name, name length, kind, colour, left and right siblings, child, class id,
# state bits, creation and modification times, first sector, size.
_DIRECTORY_ENTRY = struct.Struct("<64sHBBIII16sIQQIQ")
_UNUSED, _STREAM_OBJECT, _ROOT_STORAGE = 0, 2, 5
_RED, _BLACK = 0, 1  # the directory is a red-black tree
_STREAM_SIZE = 4096  # the smallest size kept in ordinary sectors rather than the mini stream
_WORD_97_FIB = struct.pack("<HH", 0xA5EC, 0x00C1)  # wIdent, nFib


def legacy_doc() -> bytes:
    return _compound_file((("WordDocument", _WORD_97_FIB),))


def locked_docx() -> bytes:
    """A DOCX locked with a password, as Word saves one: a compound file holding the
    encryption details and the encrypted package. Neither stream holds real content."""
    return _compound_file((("EncryptionInfo", struct.pack("<HH", 4, 4)), ("EncryptedPackage", b"")))


def docx_unpacking_to(size: int) -> bytes:
    """A minimal DOCX that unpacks to `size` bytes: its document, padded with a part of
    zeros that packs small, as a zip bomb is."""
    buffer = io.BytesIO(docx("สัญญา"))
    with zipfile.ZipFile(buffer, "a") as archive:
        left = size - sum(entry.file_size for entry in archive.infolist())
        padding = _entry("word/media/padding.bin")
        padding.compress_type = zipfile.ZIP_DEFLATED
        with archive.open(padding, "w") as part:
            while left > 0:
                chunk = min(left, 1024 * 1024)
                part.write(bytes(chunk))
                left -= chunk
    return buffer.getvalue()


def _compound_file(streams: tuple[tuple[str, bytes], ...]) -> bytes:
    """A compound file of up to three streams, each starting with the bytes given."""
    stream_sectors = _STREAM_SIZE // _SECTOR
    header = struct.pack(
        "<8s16sHHHHH6sIIIIIIIII",
        bytes.fromhex("d0cf11e0a1b11ae1"),  # signature
        bytes(16),  # class id
        0x003E,  # minor version
        0x0003,  # major version
        0xFFFE,  # byte order: little-endian
        9,  # sector size 2**9
        6,  # mini sector size 2**6
        bytes(6),
        0,  # directory sectors (always 0 in version 3)
        1,  # allocation table sectors
        1,  # first directory sector
        0,  # transaction signature
        _STREAM_SIZE,  # mini stream cutoff
        _END_OF_CHAIN,  # first mini allocation table sector: none
        0,
        _END_OF_CHAIN,  # first extra allocation table sector: none
        0,
    ) + struct.pack("<109I", 0, *[_FREE] * 108)
    # Sector 0 holds the allocation table, sector 1 the directory, then each stream in turn.
    chain = [_FAT_SECTOR, _END_OF_CHAIN]
    for index in range(len(streams)):
        first = 2 + index * stream_sectors
        chain += [first + offset + 1 for offset in range(stream_sectors - 1)] + [_END_OF_CHAIN]
    allocation_table = struct.pack(f"<{_SECTOR // 4}I", *chain, *[_FREE] * (_SECTOR // 4 - len(chain)))
    entries = [_directory_entry("Root Entry", _ROOT_STORAGE, child=1, start=_END_OF_CHAIN)]
    for index, (name, _) in enumerate(streams):
        right = index + 2 if index + 1 < len(streams) else _NO_STREAM  # the next stream, as a sibling
        entries.append(
            _directory_entry(name, _STREAM_OBJECT, right=right, start=2 + index * stream_sectors, size=_STREAM_SIZE)
        )
    entries += [_directory_entry("", _UNUSED)] * (4 - len(entries))  # four entries fill the directory sector
    contents = b"".join(start.ljust(_STREAM_SIZE, b"\0") for _, start in streams)
    return header + allocation_table + b"".join(entries) + contents


def _directory_entry(
    name: str, kind: int, *, child: int = _NO_STREAM, right: int = _NO_STREAM, start: int = 0, size: int = 0
) -> bytes:
    """An entry of the directory; an unused entry has no name."""
    encoded = (name + "\0").encode("utf-16-le") if name else b""
    colour = _BLACK if name else _RED
    return _DIRECTORY_ENTRY.pack(
        encoded, len(encoded), kind, colour, _NO_STREAM, right, child, bytes(16), 0, 0, 0, start, size
    )


@dataclass(frozen=True)
class Heading:
    text: str
    level: int = 1


@dataclass(frozen=True)
class Item:
    text: str
    level: int = 0  # 0 is the outermost


@dataclass(frozen=True)
class Numbered:
    """A list whose numbers Word draws; the text holds none of them.

    Each level has a number format (Word's numFmt, such as "decimal",
    "thaiNumbers" or "thaiLetters") and a pattern (Word's lvlText, such as "%1."
    or "%1.%2"). A list has a definition of its own, or it carries on the list
    before it, or it starts again from `start` on that list's definition, as
    Word's "Restart at 1" does. A list that starts again can redraw its first
    level in a format and pattern of its own, as a level override does.
    """

    items: tuple[str | Item, ...]
    levels: tuple[tuple[str, str], ...] = (("decimal", "%1."),)
    start: int = 1
    begins: Literal["anew", "carrying on", "starting again"] = "anew"
    redrawn: tuple[str, str] | None = None


@dataclass(frozen=True)
class NumberedHeadings:
    """Headings numbered through their styles, as Word's multilevel heading lists are:
    each heading style names the list, and each level of the list names its style.
    Word also names the level in each style below the first; `levels_in_styles` false
    leaves that out, as some other programs do."""

    headings: tuple[Heading, ...]
    levels: tuple[tuple[str, str], ...] = (("decimal", "%1"), ("decimal", "%1.%2"))
    levels_in_styles: bool = True


@dataclass(frozen=True)
class Bulleted:
    items: tuple[str | Item, ...]


@dataclass(frozen=True)
class Image:
    width: float = 120  # points
    height: float = 80


@dataclass(frozen=True)
class Table:
    rows: tuple[tuple[Cell | str | Image, ...], ...]  # an image fills a cell of its own
    # The rows that Word's last layout began a new page with.
    new_pages: tuple[int, ...] = ()
    # Rows, and cells by row and position, held in content controls, as Word's repeating sections are.
    controlled_rows: tuple[int, ...] = ()
    controlled_cells: tuple[tuple[int, int], ...] = ()


@dataclass(frozen=True)
class TextBoxes:
    """A group of text boxes drawn beside the text, each holding a line."""

    texts: tuple[str, ...]


@dataclass(frozen=True)
class WithImage:
    """A paragraph with a picture set in its text, after `before`."""

    before: str
    after: str


@dataclass(frozen=True)
class PageBreak:
    """A page break before the next block: typed into the document, set on the next
    paragraph as "page break before", or, when "rendered", where Word's last layout
    broke the page, as Word records on saving."""

    how: Literal["typed", "before", "rendered"] = "typed"


@dataclass(frozen=True)
class Split:
    """A paragraph that Word's last layout broke across two pages, after `before`."""

    before: str
    after: str


@dataclass(frozen=True)
class Deleted:
    text: str  # deleted with tracked changes on: Word shows it struck through, or not at all


@dataclass(frozen=True)
class Inserted:
    text: str  # inserted with tracked changes on


@dataclass(frozen=True)
class Link:
    text: str
    target: str = "https://example.org"


@dataclass(frozen=True)
class Hidden:
    text: str  # formatted as hidden: Word neither shows nor prints it


@dataclass(frozen=True)
class Field:
    """A field, such as a page number: the code behind it and the result Word shows."""

    code: str
    result: str


@dataclass(frozen=True)
class Rich:
    """A paragraph of plain text and the pieces Word keeps apart from it."""

    parts: tuple[str | Deleted | Inserted | Hidden | Link | Field, ...]


Block = (
    str | Heading | Numbered | NumberedHeadings | Bulleted | Table | TextBoxes | Image | WithImage | PageBreak | Split | Rich
)


@dataclass
class _Lists:
    """The numbering definitions written so far."""

    numbering: BaseOxmlElement
    last: tuple[int, int] | None = None  # the last numbered list's abstract definition and instance
    bullets: int | None = None


def document(*blocks: Block) -> bytes:
    """A DOCX of the blocks. Tabs and newlines in their text become Word's tabs and line breaks."""
    word = python_docx.Document()
    lists = _Lists(word.part.numbering_part.element)
    pending: list[str] = []  # how the breaks before the next block are made
    for block in blocks:
        if isinstance(block, PageBreak):
            if block.how == "typed":
                word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            else:
                pending.append(block.how)
            continue
        first = _add_block(word, block, lists)
        if "rendered" in pending:
            _mark_rendered_break(first)
        if "before" in pending:
            first.paragraph_format.page_break_before = True
        pending.clear()
    return _with_fixed_dates(word)


def _add_block(word: WordDocument, block: Block, lists: _Lists) -> Paragraph:
    """Add the block and return its first paragraph."""
    match block:
        case str():
            return _text(word.add_paragraph(), block)
        case Heading(text, level):
            return _text(word.add_paragraph(style=f"Heading {level}"), text)
        case Split(before, after):
            paragraph = _text(word.add_paragraph(), before)
            run = paragraph.add_run(after)
            run._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
            return paragraph
        case Numbered():
            return _items(word, block.items, _numbered_list(lists, block))
        case NumberedHeadings(headings, levels, levels_in_styles):
            styles = tuple(f"Heading{level + 1}" for level in range(len(levels)))
            _, instance = _add_definition(lists, levels, 1, styles)
            for level, style_id in enumerate(styles):
                style = word.styles.element.find(f"{qn('w:style')}[@{qn('w:styleId')}='{style_id}']")
                numbering = style.get_or_add_pPr().get_or_add_numPr()
                if level and levels_in_styles:
                    numbering.get_or_add_ilvl().val = level
                numbering.get_or_add_numId().val = instance
            paragraphs = [_text(word.add_paragraph(style=f"Heading {h.level}"), h.text) for h in headings]
            return paragraphs[0]
        case TextBoxes(texts):
            shapes = "".join(
                f"<v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p>"
                "</w:txbxContent></v:textbox></v:shape>"
                for text in texts
            )
            paragraph = word.add_paragraph()
            paragraph._p.append(
                parse_xml(
                    f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:group>{shapes}'
                    "</v:group></w:pict></w:r>"
                )
            )
            return paragraph
        case Bulleted(items):
            if lists.bullets is None:
                # Word defines bullets at every level; three are enough here.
                lists.bullets = _add_definition(lists, (("bullet", chr(0x2022)),) * 3, 1)[1]
            return _items(word, items, lists.bullets)
        case Table():
            return _table(word, block)
        case Image():
            return _picture(word.add_paragraph(), block)
        case WithImage(before, after):
            paragraph = _picture(_text(word.add_paragraph(), before), Image())
            return _text(paragraph, after)
        case Rich(parts):
            paragraph = word.add_paragraph()
            for number, part in enumerate(parts):
                paragraph._p.extend(_rich_part(word, part, number))
            return paragraph
    raise AssertionError(f"not a block of its own: {block}")


def _rich_part(
    word: WordDocument, part: str | Deleted | Inserted | Hidden | Link | Field, number: int
) -> list[BaseOxmlElement]:
    """The paragraph content Word writes for one part; `number` tells revisions apart."""

    def run(inner: str) -> str:
        return f"<w:r {nsdecls('w')}>{inner}</w:r>"

    def text(value: str, tag: str = "w:t") -> str:
        return f'<{tag} xml:space="preserve">{escape(value)}</{tag}>'

    revision = f'w:id="{number}" w:author="ผู้ตรวจ" w:date="2026-10-20T00:00:00Z"'
    match part:
        case str():
            xml = [run(text(part))]
        case Deleted(value):
            xml = [f"<w:del {nsdecls('w')} {revision}>{run(text(value, 'w:delText'))}</w:del>"]
        case Hidden(value):
            xml = [f"<w:r {nsdecls('w')}><w:rPr><w:vanish/></w:rPr>{text(value)}</w:r>"]
        case Inserted(value):
            xml = [f"<w:ins {nsdecls('w')} {revision}>{run(text(value))}</w:ins>"]
        case Link(value, target):
            relationship = word.part.relate_to(target, RT.HYPERLINK, is_external=True)
            xml = [f"<w:hyperlink {nsdecls('w', 'r')} r:id=\"{relationship}\">{run(text(value))}</w:hyperlink>"]
        case Field(code, result):
            xml = [
                run('<w:fldChar w:fldCharType="begin"/>'),
                run(f'<w:instrText xml:space="preserve"> {escape(code)} </w:instrText>'),
                run('<w:fldChar w:fldCharType="separate"/>'),
                run(text(result)),
                run('<w:fldChar w:fldCharType="end"/>'),
            ]
    return [parse_xml(fragment) for fragment in xml]


def _picture(paragraph: Paragraph, image: Image) -> Paragraph:
    paragraph.add_run().add_picture(io.BytesIO(_png()), width=Pt(image.width), height=Pt(image.height))
    return paragraph


def _text(paragraph: Paragraph, text: str) -> Paragraph:
    for index, line in enumerate(text.split("\n")):
        if index:
            paragraph.add_run().add_break()
        for piece_index, piece in enumerate(line.split("\t")):
            run = paragraph.add_run()
            if piece_index:
                run._r.append(OxmlElement("w:tab"))
            run.add_text(piece)
    return paragraph


def _items(word: WordDocument, items: tuple[str | Item, ...], number_id: int) -> Paragraph:
    """Add the items as paragraphs of the list `number_id`, and return the first."""
    paragraphs = []
    for entry in items:
        item = entry if isinstance(entry, Item) else Item(entry)
        paragraph = _text(word.add_paragraph(style="List Paragraph"), item.text)
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_ilvl().val = item.level
        numbering.get_or_add_numId().val = number_id
        paragraphs.append(paragraph)
    return paragraphs[0]


def _numbered_list(lists: _Lists, block: Numbered) -> int:
    if block.begins == "carrying on" and lists.last:
        return lists.last[1]
    if block.begins == "starting again" and lists.last:
        instance = lists.numbering.add_num(lists.last[0])
        redrawn = ""
        if block.redrawn:
            number_format, pattern = block.redrawn
            redrawn = (
                f'<w:lvl w:ilvl="0"><w:start w:val="{block.start}"/><w:numFmt w:val="{number_format}"/>'
                f'<w:lvlText w:val="{escape(pattern)}"/></w:lvl>'
            )
        instance.append(
            parse_xml(
                f'<w:lvlOverride {nsdecls("w")} w:ilvl="0"><w:startOverride w:val="{block.start}"/>{redrawn}'
                "</w:lvlOverride>"
            )
        )
        lists.last = (lists.last[0], int(instance.numId))
        return lists.last[1]
    lists.last = _add_definition(lists, block.levels, block.start)
    return lists.last[1]


def _add_definition(
    lists: _Lists, levels: tuple[tuple[str, str], ...], start: int, styles: tuple[str, ...] = ()
) -> tuple[int, int]:
    """A new abstract numbering definition and an instance of it: their ids. A level may name
    the paragraph style it numbers, as Word's heading lists do."""
    taken = [int(abstract.get(qn("w:abstractNumId"))) for abstract in lists.numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(taken) + 1
    levels_xml = "".join(
        f'<w:lvl w:ilvl="{level}"><w:start w:val="{start}"/><w:numFmt w:val="{number_format}"/>'
        + (f'<w:pStyle w:val="{styles[level]}"/>' if level < len(styles) else "")
        + f'<w:lvlText w:val="{escape(pattern)}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="{720 * (level + 1)}" w:hanging="360"/></w:pPr></w:lvl>'
        for level, (number_format, pattern) in enumerate(levels)
    )
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="multilevel"/>{levels_xml}</w:abstractNum>'
    )
    # Every abstract definition comes before the first instance.
    lists.numbering.find(qn("w:num")).addprevious(abstract)
    return abstract_id, int(lists.numbering.add_num(abstract_id).numId)


def _table(word: WordDocument, block: Table) -> Paragraph:
    """Add a table with its merged cells, and return the first cell's paragraph."""
    rows = block.rows
    placed = []
    covered: set[tuple[int, int]] = set()
    for row_index, row in enumerate(rows):
        column = 0
        for entry in row:
            cell = entry if isinstance(entry, Cell) else Cell("") if isinstance(entry, Image) else Cell(entry)
            while (row_index, column) in covered:
                column += 1
            placed.append((row_index, column, cell, entry if isinstance(entry, Image) else None))
            covered.update(
                (row_index + down, column + across) for down in range(cell.rowspan) for across in range(cell.colspan)
            )
            column += cell.colspan
    table = word.add_table(rows=len(rows), cols=max(column for _, column in covered) + 1)
    table.style = "Table Grid"
    for row_index, column, cell, image in placed:
        target = table.cell(row_index, column)
        if cell.rowspan > 1 or cell.colspan > 1:
            target = target.merge(table.cell(row_index + cell.rowspan - 1, column + cell.colspan - 1))
        _text(target.paragraphs[0], cell.text)
        if image:
            _picture(target.paragraphs[0], image)
    for row_index in block.new_pages:
        _mark_rendered_break(table.cell(row_index, 0).paragraphs[0])
    first: Paragraph = table.cell(0, 0).paragraphs[0]
    row_elements = table._tbl.findall(qn("w:tr"))
    for row_index, column in block.controlled_cells:
        _in_content_control(row_elements[row_index].findall(qn("w:tc"))[column])
    for row_index in block.controlled_rows:
        _in_content_control(row_elements[row_index])
    return first


def _in_content_control(element: BaseOxmlElement) -> None:
    control = parse_xml(f'<w:sdt {nsdecls("w")}><w:sdtPr/><w:sdtContent/></w:sdt>')
    element.addprevious(control)
    control.find(qn("w:sdtContent")).append(element)


def _mark_rendered_break(paragraph: Paragraph) -> None:
    """Word records where its layout broke the page in the run that begins the new page."""
    runs = paragraph._p.findall(qn("w:r"))
    run = runs[0] if runs else paragraph.add_run()._r
    run.insert(0 if run.rPr is None else 1, OxmlElement("w:lastRenderedPageBreak"))


def _png() -> bytes:
    pixels = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 24, 16), False)
    pixels.set_rect(pixels.irect, (90, 140, 200))
    png: bytes = pixels.tobytes("png")
    return png


def _with_fixed_dates(word: WordDocument) -> bytes:
    """The saved file, its zip entries written alike, so the same blocks always give the same bytes."""
    saved = io.BytesIO()
    word.save(saved)
    fixed = io.BytesIO()
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(fixed, "w") as target:
        for entry in source.infolist():
            target.writestr(_entry(entry.filename), source.read(entry.filename), zipfile.ZIP_DEFLATED)
    return fixed.getvalue()
