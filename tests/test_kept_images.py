"""Kept images, through the conversion function: scans kept as images of the page instead of read
with OCR, and pictures on text pages shown as themselves in the DOCX."""

import io
from pathlib import Path
from typing import Literal

import docx
import markdown_reader
import numpy as np
import pymupdf
import pytest
from docx_reader import A4, InlineImage, Layout, block_text, inline_images, page_texts, read_docx
from fake_ocr import FakeEngine
from generate import CLAUSES, SIGNATURE, SIGNATURE_LINE
from pdf_builder import TOP, Drawn, Grid, Line, Pasted, Picture, Rotated, Scan, Styled, highlighted, pdf, remapped
from word_builder import PageBreak, document

from pdf_to_doc import ConversionResult, Opened, Progress, convert
from pdf_to_doc.cli import main
from pdf_to_doc.intake import MAX_OUTPUTS
from pdf_to_doc.model import LARGEST_PAGE

SARA_AA = chr(0x0E32)
SARA_AM = chr(0x0E33)
# A page as pdf_builder makes it, and as a kept page fills its Word page from edge to edge.
A4_EDGE_TO_EDGE = Layout(595.0, 842.0, False, (0.0, 0.0, 0.0, 0.0))


def markdown_pages(result: ConversionResult) -> list[str]:
    return markdown_reader.page_texts(result.outputs["markdown"].decode("utf-8"))


def keep_images(source: bytes) -> ConversionResult:
    return convert(source, to={"markdown", "docx"}, ocr=FakeEngine(), keep_images=True)


def kept_page_image(data: bytes, page: int = 0) -> InlineImage:
    """The image a page of a DOCX output holds alone, as a kept page does."""
    [block] = read_docx(data).pages[page]
    [image] = inline_images(block)
    return image


def resolution(image: InlineImage) -> tuple[int, int]:
    """The image's pixels per inch, across and down, as shown."""
    pixels = pymupdf.Pixmap(image.data)
    return round(pixels.width / image.width * 72), round(pixels.height / image.height * 72)


def ink_extent(image: InlineImage) -> tuple[float, float, float, float]:
    """Where the image's dark pixels lie: their left, top, right and bottom, as shares of its width and height."""
    grey = pymupdf.Pixmap(pymupdf.csGRAY, pymupdf.Pixmap(image.data))
    pixels = np.frombuffer(grey.samples, dtype=np.uint8).reshape(grey.height, grey.stride)[:, : grey.width]
    rows, columns = np.nonzero(pixels < 128)
    return columns.min() / grey.width, rows.min() / grey.height, columns.max() / grey.width, rows.max() / grey.height


def test_a_scan_is_kept_never_read_with_ocr_and_stands_as_a_placeholder_in_markdown() -> None:
    engine = FakeEngine()

    result = convert(pdf(Scan(["ข้อความจากภาพสแกน"], dpi=50)), ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["[หน้าสแกน]"]
    assert [page.kind for page in result.summary.pages] == ["kept"]
    assert engine.images == []


def test_a_scan_pasted_over_a_third_of_a_page_is_still_a_scan() -> None:
    # What the page shows lies in the scan, though the scan covers only part of the page.
    result = keep_images(pdf(Pasted(595, 842, box=(72, 150, 523, 400))))

    assert [page.kind for page in result.summary.pages] == ["kept"]


def test_a_page_showing_nothing_but_a_small_picture_is_kept_too() -> None:
    # Decided with the user on 2026-10-24: its image is all there is to keep, and OCR
    # would read the picture as nonsense.
    result = keep_images(pdf(Pasted(595, 842, box=(272, 400, 322, 450))))

    assert [page.kind for page in result.summary.pages] == ["kept"]


# DOCX: a kept page is an image of the whole page, from edge to edge


def test_a_kept_scan_fills_a_word_page_of_its_own_from_edge_to_edge_at_150_dpi() -> None:
    result = keep_images(pdf(["หน้าที่มีข้อความ"], Scan(["หน้าสแกน"], dpi=50), ["หน้าสุดท้าย"]))

    data = result.outputs["docx"]

    assert read_docx(data).layouts == [A4, A4_EDGE_TO_EDGE, A4]
    assert page_texts(data) == ["หน้าที่มีข้อความ", "", "หน้าสุดท้าย"]
    image = kept_page_image(data, page=1)
    assert (image.width, image.height) == (595.0, 842.0)
    assert resolution(image) == (150, 150)


def test_kept_pages_first_last_and_one_after_another_each_take_a_word_page_with_nothing_left_over() -> None:
    source = pdf(Scan(["สแกนแรก"], dpi=50), Scan(["สแกนที่สอง"], dpi=50), ["หน้าที่มีข้อความ"], Scan(["สแกนสุดท้าย"], dpi=50))

    data = keep_images(source).outputs["docx"]

    output = read_docx(data)
    # With page 1 kept, the summary has the first Word page, on A4, to itself.
    assert output.summary_layout == A4
    assert output.layouts == [A4_EDGE_TO_EDGE, A4_EDGE_TO_EDGE, A4, A4_EDGE_TO_EDGE]
    assert [[len(inline_images(block)) for block in page] for page in output.pages] == [[1], [1], [0], [1]]
    # An empty paragraph after an image as tall as its page would fall onto a page of its own.
    word = docx.Document(io.BytesIO(data))
    assert [paragraph for paragraph in word.paragraphs if not paragraph.text and not inline_images(paragraph)] == []


def test_a_kept_page_after_a_page_ending_in_a_table_still_takes_a_word_page_of_its_own() -> None:
    grid = Grid(rows=(("งวดที่", "จำนวนเงิน"), ("๑", "50,000")), widths=(150, 150))

    output = read_docx(keep_images(pdf(["ตารางงวดเงิน", grid], Scan(["สแกน"], dpi=50))).outputs["docx"])

    assert output.layouts == [A4, A4_EDGE_TO_EDGE]
    assert [[type(block).__name__ for block in page] for page in output.pages] == [["Paragraph", "Table"], ["Paragraph"]]


def test_an_empty_page_after_a_kept_page_is_still_a_page() -> None:
    data = keep_images(pdf(Scan(["สแกน"], dpi=50), [], ["หน้าสุดท้าย"])).outputs["docx"]

    assert read_docx(data).layouts == [A4_EDGE_TO_EDGE, A4, A4]
    assert page_texts(data) == ["", "", "หน้าสุดท้าย"]


def test_a_landscape_scan_keeps_a_landscape_word_page() -> None:
    data = keep_images(pdf(Scan(["ตารางแนวนอน"], dpi=50, landscape=True))).outputs["docx"]

    assert read_docx(data).layouts == [Layout(842.0, 595.0, True, (0.0, 0.0, 0.0, 0.0))]
    assert resolution(kept_page_image(data)) == (150, 150)


@pytest.mark.parametrize("rotation", [90, 270])
def test_a_scan_turned_upright_by_the_pdf_is_kept_upright(rotation: Literal[90, 270]) -> None:
    data = keep_images(pdf(Rotated(Scan(["สแกนที่เก็บไว้ตะแคง"], dpi=50), rotation))).outputs["docx"]

    assert read_docx(data).layouts == [A4_EDGE_TO_EDGE]
    left, top, right, bottom = ink_extent(kept_page_image(data))
    # Its line reads across the top left of the page, where it was printed, not down a side.
    assert left < 0.2 and top < 0.2
    assert (right - left) * 595 > (bottom - top) * 842


def test_a_sideways_scan_is_kept_as_it_is() -> None:
    data = keep_images(pdf(Scan(["ตารางที่พิมพ์ตะแคง"], dpi=50, sideways=90))).outputs["docx"]

    assert read_docx(data).layouts == [A4_EDGE_TO_EDGE]
    left, top, right, bottom = ink_extent(kept_page_image(data))
    # Its line still runs down the page, as in the original.
    assert (bottom - top) * 842 > (right - left) * 595


def test_a_page_larger_than_word_takes_is_scaled_down_keeping_its_shape_and_150_dpi_at_that_size() -> None:
    # 2000 by 1000 points, about 28 by 14 inches; Word takes 22 inches at most.
    data = keep_images(pdf(Pasted(2000, 1000))).outputs["docx"]

    assert read_docx(data).layouts == [Layout(1584.0, 792.0, True, (0.0, 0.0, 0.0, 0.0))]
    image = kept_page_image(data)
    assert (image.width, image.height, resolution(image)) == (1584.0, 792.0, (150, 150))


def test_a_highlight_on_a_scan_shows_in_its_kept_image() -> None:
    source = highlighted(pdf(Scan(CLAUSES, dpi=50)), (100, 300, 300, 340))

    image = kept_page_image(keep_images(source).outputs["docx"])

    red, green, blue = pymupdf.Pixmap(image.data).pixel(round(200 * 150 / 72), round(320 * 150 / 72))
    assert red > 200 and green > 200 and blue < 100


# Markdown: the placeholder, then the text the file holds on the scan


def test_a_note_typed_over_a_scan_follows_its_placeholder_and_a_header_over_every_scan_is_left_out() -> None:
    header = Styled("เอกสารปกปิด ห้ามเผยแพร่ก่อนได้รับอนุญาต", size=14, y=40)
    note = Styled("หมายเหตุ แนบสำเนาบัตรประชาชน", size=12, y=600)
    source = pdf(Scan(CLAUSES, overlay=[header, note], dpi=50), Scan(CLAUSES, overlay=[header], dpi=50))

    result = keep_images(source)

    assert markdown_pages(result) == ["[หน้าสแกน]\n\nหมายเหตุ แนบสำเนาบัตรประชาชน", "[หน้าสแกน]"]
    assert [page.kind for page in result.summary.pages] == ["kept", "kept"]
    # The DOCX shows the note in the page's image alone.
    assert page_texts(result.outputs["docx"]) == ["", ""]


def test_large_text_on_a_kept_page_makes_no_heading() -> None:
    title = Styled("ภาคผนวกแนบท้ายสัญญา", size=24, y=300)

    result = keep_images(pdf(["ข้อความของสัญญา"], Scan(CLAUSES, overlay=[title], dpi=50)))

    assert markdown_pages(result) == ["ข้อความของสัญญา", "[หน้าสแกน]\n\nภาคผนวกแนบท้ายสัญญา"]


# Scans made searchable: kept, their hidden text following the placeholder


@pytest.mark.parametrize("hidden", ["over", "under"])
def test_a_searchable_scan_is_kept_and_its_hidden_text_follows_its_placeholder(
    hidden: Literal["over", "under"],
) -> None:
    engine = FakeEngine()

    result = convert(pdf(Scan(CLAUSES, dpi=50, searchable=hidden)), ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["[หน้าสแกน]\n\n" + "".join(CLAUSES)]
    assert [page.kind for page in result.summary.pages] == ["kept"]
    assert engine.images == []


@pytest.mark.parametrize("hidden", ["over", "under"])
def test_a_searchable_scan_turned_upright_by_the_pdf_is_kept_with_its_hidden_text(
    hidden: Literal["over", "under"],
) -> None:
    result = keep_images(pdf(Rotated(Scan(CLAUSES, dpi=50, searchable=hidden), 90)))

    assert markdown_pages(result) == ["[หน้าสแกน]\n\n" + "".join(CLAUSES)]


def test_hidden_text_makes_no_heading_and_a_hidden_header_over_every_scan_is_left_out() -> None:
    header = Styled("เอกสารปกปิด ห้ามเผยแพร่ก่อนได้รับอนุญาต", size=14, y=40)
    title = Styled("ภาคผนวกแนบท้ายสัญญา", size=24, y=150)
    scan = Scan([header, title, *CLAUSES], dpi=50, searchable="over")

    result = keep_images(pdf(["ข้อความของสัญญา"], scan, scan))

    kept = "[หน้าสแกน]\n\nภาคผนวกแนบท้ายสัญญา\n\n" + "".join(CLAUSES)
    assert markdown_pages(result) == ["ข้อความของสัญญา", kept, kept]


# What a reader sees is not hidden


def test_visible_text_over_a_letterhead_image_is_still_read_from_the_text_layer() -> None:
    engine = FakeEngine()

    result = convert(pdf(Scan(["หัวจดหมาย"], overlay=CLAUSES, dpi=50)), ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["".join(CLAUSES)]
    assert [page.kind for page in result.summary.pages] == ["text"]
    assert engine.images == []


@pytest.mark.parametrize("letterhead", [False, True], ids=["on a plain page", "on a letterhead image"])
def test_a_transparent_picture_over_text_does_not_hide_it(letterhead: bool) -> None:
    # A stamp on a transparent ground, pressed over all the text, which still shows through. On a
    # letterhead it lies within the opaque image behind the text, which must not count for it.
    stamp = Picture("stamp", 350, 350, y=TOP - 20)
    page: Scan | list[Line] = Scan(["หัวจดหมาย"], overlay=[*CLAUSES, stamp], dpi=50) if letterhead else [*CLAUSES, stamp]

    result = keep_images(pdf(page))

    assert [page.kind for page in result.summary.pages] == ["text"]
    assert markdown_pages(result) == ["[ตราประทับ]\n\n" + "".join(CLAUSES)]


# A scan whose text layer cannot be trusted is kept without it


def test_a_searchable_scan_whose_hidden_text_is_wrong_is_kept_without_that_text_and_the_summary_says_so() -> None:
    engine = FakeEngine()
    # The scanner's text reports sara aa as sara am, so "การ" comes out as "กำร".
    source = remapped(pdf(Scan(CLAUSES, dpi=50, searchable="over")), {SARA_AA: SARA_AM})

    result = convert(source, to={"markdown", "docx"}, ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["[หน้าสแกน]"]
    assert [(page.kind, page.untrusted_text_layer) for page in result.summary.pages] == [("kept", True)]
    assert engine.images == []
    assert result.summary.headline() == (
        "1 หน้า: อ่านจากข้อความในไฟล์ 0 หน้า, อ่านด้วย OCR 0 หน้า, เก็บเป็นภาพ 1 หน้า ได้แก่หน้า 1"
        " (หน้า 1 ข้อความในไฟล์ไม่น่าเชื่อถือ จึงไม่ได้นำมาใช้)"
    )


def test_a_scan_under_a_header_whose_text_is_wrong_is_kept_without_the_header() -> None:
    engine = FakeEngine()
    header = Styled("เอกสารปกปิด ห้ามเผยแพร่ก่อนได้รับอนุญาต", size=14, y=40)
    # The header's font reports sara aa as a character it cannot map.
    source = remapped(pdf(Scan(CLAUSES, overlay=[header], dpi=50)), {SARA_AA: chr(0xFFFD)})

    result = convert(source, ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["[หน้าสแกน]"]
    assert [page.kind for page in result.summary.pages] == ["kept"]
    assert engine.images == []


# Without kept images, a scan made searchable is read from its text layer


@pytest.mark.parametrize(
    ("hidden", "wrong", "kind"),
    [("over", False, "text"), ("under", False, "text"), ("over", True, "ocr")],
    ids=["hidden over", "hidden under", "hidden text wrong"],
)
def test_without_kept_images_a_searchable_scan_is_read_from_its_text_layer_unless_that_is_wrong(
    hidden: Literal["over", "under"], wrong: bool, kind: str
) -> None:
    source = pdf(Scan(CLAUSES, dpi=50, searchable=hidden))

    result = convert(remapped(source, {SARA_AA: SARA_AM}) if wrong else source, ocr=FakeEngine(CLAUSES))

    assert markdown_pages(result) == ["".join(CLAUSES)]
    assert [page.kind for page in result.summary.pages] == [kind]


def test_a_docx_source_is_converted_as_without_kept_images() -> None:
    source = document("หน้าแรก", PageBreak(), "หน้าที่สอง")

    assert convert(source, keep_images=True).output == convert(source).output


# Pictures on text pages: shown as themselves in the DOCX, their placeholders kept in the Markdown

# Enough body text for a text page around a picture.
BEFORE, AFTER = CLAUSES[:3], CLAUSES[3:]


def docx_shown(data: bytes) -> list[list[str | tuple[float, float]]]:
    """Each page of a DOCX: each paragraph's text, or the size of the one image it holds."""
    shown: list[list[str | tuple[float, float]]] = []
    for page in read_docx(data).pages:
        shown.append([])
        for block in page:
            images = inline_images(block)
            shown[-1].append((images[0].width, images[0].height) if images else block_text(block))
    return shown


def pictures_in(data: bytes) -> list[InlineImage]:
    return [image for page in read_docx(data).pages for block in page for image in inline_images(block)]


def red_slopes(pixels: pymupdf.Pixmap) -> tuple[float, float]:
    """How much redder an RGB image is in its bottom half than its top, and in its right half than its left."""
    red = np.frombuffer(pixels.samples, dtype=np.uint8).reshape(pixels.height, pixels.stride)[:, 0 : pixels.width * pixels.n : pixels.n]
    height, width = red.shape
    return red[height // 2 :].mean() - red[: height // 2].mean(), red[:, width // 2 :].mean() - red[:, : width // 2].mean()


def test_a_photo_between_paragraphs_is_shown_as_itself_where_its_placeholder_was() -> None:
    source = pdf([*BEFORE, Picture("photo", 120, 80), *AFTER])

    result = keep_images(source)

    assert docx_shown(result.outputs["docx"]) == [["".join(BEFORE), (120.0, 80.0), "".join(AFTER)]]
    assert markdown_pages(result) == ["".join(BEFORE) + "\n\n[รูปภาพ]\n\n" + "".join(AFTER)]
    [photo] = pictures_in(result.outputs["docx"])
    assert photo.data.startswith(b"\xff\xd8")  # a JPEG: an opaque photo has no transparency to keep


def test_a_signature_on_its_signature_line_is_shown_as_itself_on_a_transparent_ground() -> None:
    data = keep_images(pdf([*CLAUSES, "", SIGNATURE, SIGNATURE_LINE])).outputs["docx"]

    assert docx_shown(data) == [["".join(CLAUSES), (140.0, 45.0), SIGNATURE_LINE]]
    [signature] = pictures_in(data)
    assert pymupdf.Pixmap(signature.data).alpha  # the line under it shows through


@pytest.mark.parametrize("kind", ["stamp", "chart"])
def test_a_stamp_or_a_chart_is_shown_as_itself_on_a_transparent_ground(kind: Literal["stamp", "chart"]) -> None:
    data = keep_images(pdf([*BEFORE, Picture(kind, 70, 70), *AFTER])).outputs["docx"]

    assert docx_shown(data) == [["".join(BEFORE), (70.0, 70.0), "".join(AFTER)]]
    [picture] = pictures_in(data)
    assert pymupdf.Pixmap(picture.data).alpha


def faces_as_on_the_page(source: bytes, box: pymupdf.Rect) -> None:
    """Assert that the one picture kept from the source, a photo growing redder along its own width,
    lies as the photo in the box does on the source's first page, rendered as shown."""
    [photo] = pictures_in(keep_images(source).outputs["docx"])
    on_the_page = pymupdf.open(stream=source, filetype="pdf")[0].get_pixmap(clip=box, colorspace=pymupdf.csRGB)
    assert (photo.width, photo.height) == (box.width, box.height)
    # Counted in steps of 20 grey levels: the photo's red runs from 30 to 230 across it, so a half
    # that faces the other way differs by about 100, and noise by a few.
    kept, shown = red_slopes(pymupdf.Pixmap(photo.data)), red_slopes(on_the_page)
    assert [round(slope / 20) for slope in kept] == [round(slope / 20) for slope in shown]


@pytest.mark.parametrize("rotation", [90, 270])
def test_a_picture_on_a_page_turned_upright_by_the_pdf_stands_upright(rotation: Literal[90, 270]) -> None:
    box = pymupdf.Rect(72, 300, 192, 380)

    faces_as_on_the_page(pdf(Rotated([*CLAUSES, Picture("photo", box.width, box.height, y=box.y0)], rotation)), box)


def test_a_signature_on_a_page_turned_upright_by_the_pdf_keeps_its_transparency() -> None:
    [signature] = pictures_in(keep_images(pdf(Rotated([*CLAUSES, "", SIGNATURE, SIGNATURE_LINE], 90))).outputs["docx"])

    assert (signature.width, signature.height) == (140.0, 45.0)
    assert pymupdf.Pixmap(signature.data).alpha


@pytest.mark.parametrize("turn", [90, 180, 270])
def test_a_picture_drawn_turned_is_shown_turned_as_on_the_page(turn: Literal[90, 180, 270]) -> None:
    box = pymupdf.Rect(72, 300, 192, 380) if turn == 180 else pymupdf.Rect(72, 300, 152, 420)

    faces_as_on_the_page(pdf([*CLAUSES, Picture("photo", box.width, box.height, y=box.y0, turn=turn)]), box)


def test_a_picture_sharper_than_150_dpi_is_kept_at_150_dpi_and_a_softer_one_as_it_is() -> None:
    data = keep_images(pdf([*BEFORE, Picture("photo", 120, 80, dpi=600), Picture("photo", 120, 80), *AFTER])).outputs["docx"]

    assert [resolution(picture) for picture in pictures_in(data)] == [(150, 150), (144, 144)]


def test_a_cmyk_photo_is_kept_in_rgb() -> None:
    [photo] = pictures_in(keep_images(pdf([*BEFORE, Picture("photo", 120, 80, cmyk=True), *AFTER])).outputs["docx"])

    assert pymupdf.Pixmap(photo.data).n == 3


def test_a_picture_wider_than_the_text_area_is_shrunk_to_it_keeping_its_shape() -> None:
    data = keep_images(pdf([*BEFORE, Picture("photo", 520, 100), *AFTER])).outputs["docx"]

    # The text area of A4 with margins of 2.54 cm is 451.3 points wide.
    assert docx_shown(data) == [["".join(BEFORE), (451.3, 86.8), "".join(AFTER)]]


def test_a_picture_in_a_table_cell_comes_right_after_the_table() -> None:
    grid = Grid(rows=(("ตำแหน่ง", "ลายมือ"), ("ผู้ให้ทุน", "")), widths=(200, 200), y=200)
    in_cell = Picture("photo", 60, 20, x=290, y=214)

    [blocks] = docx_shown(keep_images(pdf([*CLAUSES, grid, in_cell, Styled("ทั้งนี้ตามที่ตกลงกัน", y=320)])).outputs["docx"])

    assert blocks[2:] == [(60.0, 20.0), "ทั้งนี้ตามที่ตกลงกัน"]


def test_no_picture_is_shown_for_page_furniture_a_letterhead_behind_text_or_a_rule() -> None:
    logo = Picture("photo", 60, 30, y=30)
    body: list[Line] = [Styled(CLAUSES[0], y=160), *CLAUSES[1:]]
    rule = Picture("photo", 400, 2)
    source = pdf([logo, *body], [logo, *body, rule], Scan(["หัวจดหมาย"], overlay=CLAUSES, dpi=50))

    assert pictures_in(keep_images(source).outputs["docx"]) == []


def test_the_markdown_of_pictures_is_the_same_with_kept_images() -> None:
    source = pdf([*CLAUSES, Picture("photo", 120, 80), "", SIGNATURE, SIGNATURE_LINE, "", Picture("stamp", 70, 70)])

    assert keep_images(source).outputs["markdown"] == convert(source, to="markdown", ocr=FakeEngine()).output


# Outputs within 50 MB


@pytest.mark.slow
def test_outputs_that_would_pass_50_mb_fit_once_their_images_are_lowered() -> None:
    # Colour noise is as hard as a scan gets to compress: seven pages of it, 22 by 22 inches,
    # come to about 52 MB of page images at 150 dpi. Compressing them takes most of a minute.
    source = pdf(*(Pasted(LARGEST_PAGE, LARGEST_PAGE, noise_pixels=800, seed=seed) for seed in range(7)))

    result = keep_images(source)

    assert sum(len(output) for output in result.outputs.values()) <= MAX_OUTPUTS
    lowered_to = result.summary.images_lowered_to
    assert lowered_to is not None and lowered_to < 150
    assert f"ภาพใน DOCX ลดความละเอียดเหลือ {lowered_to} dpi ให้ผลลัพธ์ไม่เกิน 50 MB" in result.summary.headline()
    assert all(max(resolution(image)) <= lowered_to for image in pictures_in(result.outputs["docx"]))


def test_outputs_within_50_mb_keep_their_images_at_150_dpi_and_say_nothing_of_it() -> None:
    result = keep_images(pdf(Scan(CLAUSES, dpi=50), [*BEFORE, Picture("photo", 120, 80), *AFTER]))

    assert result.summary.images_lowered_to is None
    assert "dpi" not in result.summary.headline()
    assert [max(resolution(image)) for image in pictures_in(result.outputs["docx"])] == [150, 144]


# The summary


def test_the_summary_names_the_kept_pages_and_counts_no_placeholder_for_them() -> None:
    source = pdf(["หน้าที่มีข้อความ"], Scan(["สแกน"], dpi=50), Scan(["สแกน"], dpi=50), ["ข้อความ"], Scan(["สแกน"], dpi=50))

    result = keep_images(source)

    headline = result.summary.headline()
    assert headline == "5 หน้า: อ่านจากข้อความในไฟล์ 2 หน้า, อ่านด้วย OCR 0 หน้า, เก็บเป็นภาพ 3 หน้า ได้แก่หน้า 2–3, 5"
    assert headline in result.outputs["markdown"].decode("utf-8")
    assert headline in read_docx(result.outputs["docx"]).summary.text


def test_a_file_with_nothing_to_keep_reports_and_sums_up_as_without_kept_images() -> None:
    source = pdf(["หน้าที่มีข้อความ"], Drawn(["พิมพ์เป็นเส้น"]))
    reports: dict[bool, list[Opened | Progress]] = {False: [], True: []}

    results = {
        keep: convert(source, ocr=FakeEngine(["พิมพ์เป็นเส้น"]), keep_images=keep, progress=reports[keep].append)
        for keep in (False, True)
    }

    assert reports[True] == reports[False]
    assert results[True].output == results[False].output


# Progress and the command line


def test_each_page_image_is_reported_as_current_of_total_before_the_ocr_pages() -> None:
    reports: list[Opened | Progress] = []

    convert(
        pdf(Scan(["สแกน"], dpi=50), Drawn(["พิมพ์เป็นเส้น"]), Scan(["สแกน"], dpi=50)),
        ocr=FakeEngine(["พิมพ์เป็นเส้น"]),
        keep_images=True,
        progress=reports.append,
    )

    assert reports == [
        Opened(source_type="pdf", pages=3),
        Progress(stage="reading"),
        Progress(stage="images", current=1, total=2),
        Progress(stage="images", current=2, total=2),
        Progress(stage="ocr", current=1, total=1),
        Progress(stage="writing"),
    ]


def test_the_command_keeps_images_and_shows_progress_while_it_makes_them(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "scans.pdf"
    source.write_bytes(pdf(Scan(["หน้าสแกนแรก"], dpi=50), Scan(["หน้าสแกนที่สอง"], dpi=50)))

    assert main([str(source), "--keep-images", "--to", "docx"], ocr=FakeEngine()) == 0

    shown = capsys.readouterr().err
    assert "เก็บภาพหน้า 1/2" in shown
    assert "เก็บภาพหน้า 2/2" in shown
    assert read_docx((tmp_path / "scans.docx").read_bytes()).layouts == [A4_EDGE_TO_EDGE, A4_EDGE_TO_EDGE]


def test_a_file_of_scans_kept_as_images_needs_no_ocr_models(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PDF_TO_DOC_MODELS", str(tmp_path / "no-models-here"))
    source = tmp_path / "scan.pdf"
    source.write_bytes(pdf(Scan(["หน้าสแกน"], dpi=50)))

    assert main([str(source), "--keep-images"]) == 0

    assert markdown_reader.page_texts((tmp_path / "scan.md").read_text(encoding="utf-8")) == ["[หน้าสแกน]"]


# What is not a scan is read as without kept images


def test_a_page_whose_text_is_drawn_as_shapes_is_still_read_with_ocr() -> None:
    engine = FakeEngine(["ข้อความที่พิมพ์ออกมาเป็นเส้น"])

    result = convert(pdf(Drawn(["ข้อความที่พิมพ์ออกมาเป็นเส้น"])), ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["ข้อความที่พิมพ์ออกมาเป็นเส้น"]
    assert [page.kind for page in result.summary.pages] == ["ocr"]
    assert len(engine.images) == 1


def test_a_page_whose_text_layer_cannot_be_trusted_is_still_read_with_ocr() -> None:
    # The font reports sara aa as sara am, so "การ" comes out as "กำร".
    source = remapped(pdf(CLAUSES), {SARA_AA: SARA_AM})

    result = convert(source, ocr=FakeEngine(CLAUSES), keep_images=True)

    assert markdown_pages(result) == ["".join(CLAUSES)]
    assert [(page.kind, page.untrusted_text_layer) for page in result.summary.pages] == [("ocr", True)]


def test_a_text_page_is_read_from_its_text_layer() -> None:
    engine = FakeEngine()

    result = convert(pdf(["ข้อความในไฟล์"], Scan(["หน้าสแกน"], dpi=50)), ocr=engine, keep_images=True)

    assert markdown_pages(result) == ["ข้อความในไฟล์", "[หน้าสแกน]"]
    assert [page.kind for page in result.summary.pages] == ["text", "kept"]
    assert engine.images == []
