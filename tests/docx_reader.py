"""Reads a DOCX output back the way Word shows it: the summary, then page by page, how
each page is laid out, the images in it, and how Word draws each run of text."""

import io
from dataclasses import dataclass, replace
from typing import cast

import docx
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.parts.document import DocumentPart
from docx.section import Section
from docx.shared import Emu
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

Block = Paragraph | Table


@dataclass(frozen=True)
class Layout:
    """A Word page's size and margins, in points rounded to a tenth."""

    width: float
    height: float
    landscape: bool
    margins: tuple[float, float, float, float]  # top, right, bottom, left


A4 = Layout(595.3, 841.9, False, (72.0, 72.0, 72.0, 72.0))  # with Word's usual margins of 2.54 cm


@dataclass(frozen=True)
class InlineImage:
    """An image set in a paragraph's line, as Word stores and shows it."""

    data: bytes  # the image file
    width: float  # in points rounded to a tenth
    height: float


@dataclass(frozen=True)
class Output:
    summary: Paragraph  # the first paragraph
    summary_layout: Layout  # how the page the summary opens is laid out
    pages: list[list[Block]]  # the paragraphs and tables of each page, without empty paragraphs
    layouts: list[Layout]  # how each page is laid out


@dataclass(frozen=True)
class Run:
    """A run of text and the formatting Word draws it with, after styles and defaults.

    Word draws Thai with the complex-script font, size and bold, and other text with
    the plain ones. It breaks Thai lines at word boundaries only in a run marked
    complex script, and proofs every character of such a run as Thai.
    """

    text: str
    complex_script: bool
    font: str
    size: float
    bold: bool
    thai_font: str
    thai_size: float
    thai_bold: bool
    highlight: str | None  # the highlight colour, such as "yellow"


@dataclass(frozen=True)
class Cell:
    text: str
    rowspan: int = 1
    colspan: int = 1


def read_docx(data: bytes) -> Output:
    """A page starts at a paragraph set to break the page before it, or at the start of a section,
    except the section after the summary's, when the summary has a page of its own."""
    word = docx.Document(io.BytesIO(data))
    summary, *body = word.iter_inner_content()
    assert isinstance(summary, Paragraph)
    layouts = [_layout(section) for section in word.sections]
    section = 1 if _ends_section(summary) else 0
    pages: list[list[Block]] = [[]]
    page_layouts = [layouts[section]]
    after_section = False
    for block in body:
        if _starts_page(block) or after_section:
            pages.append([])
            page_layouts.append(layouts[section])
        if isinstance(block, Table) or block.text or inline_images(block):
            pages[-1].append(block)
        after_section = _ends_section(block)
        section += after_section
    return Output(summary=summary, summary_layout=layouts[0], pages=pages, layouts=page_layouts)


def inline_images(block: Block) -> list[InlineImage]:
    """The images in a paragraph's lines, in order."""
    if isinstance(block, Table):
        return []
    found = []
    for inline in block._p.iter(qn("wp:inline")):
        extent = inline.find(qn("wp:extent"))
        [blip] = inline.iter(qn("a:blip"))
        image = block.part.related_parts[blip.get(qn("r:embed"))]
        width, height = (round(Emu(int(extent.get(side))).pt, 1) for side in ("cx", "cy"))
        found.append(InlineImage(image.blob, width, height))
    return found


def _ends_section(block: Block) -> bool:
    return isinstance(block, Paragraph) and block._p.pPr is not None and block._p.pPr.sectPr is not None


def _layout(section: Section) -> Layout:
    def points(length: object) -> float:
        return round(float(cast(Emu, length).pt), 1)

    margins = (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin)
    return Layout(
        width=points(section.page_width),
        height=points(section.page_height),
        landscape=section.orientation == WD_ORIENT.LANDSCAPE,
        margins=cast(tuple[float, float, float, float], tuple(points(margin) for margin in margins)),
    )


def page_texts(data: bytes) -> list[str]:
    """The text of each page: each paragraph, and each cell of a table, on a line of its own."""
    return ["\n".join(_lines(block) for block in page) for page in read_docx(data).pages]


def cells(table: Table) -> list[list[Cell]]:
    """The cells that start in each row, with the rows and columns they span, as Word merges them."""
    rows: list[list[Cell]] = []
    # For each grid column, where the cell starts that a vertical merge would continue.
    starts: dict[int, tuple[int, int]] = {}
    for row_index, row in enumerate(table._tbl.tr_lst):
        starting: list[Cell] = []
        column = 0
        for tc in row.tc_lst:
            if tc.vMerge == "continue":
                start_row, start_index = starts[column]
                above = rows[start_row][start_index]
                rows[start_row][start_index] = replace(above, rowspan=above.rowspan + 1)
            else:
                starts[column] = (row_index, len(starting))
                starting.append(Cell(_Cell(tc, table).text, colspan=tc.grid_span))
            column += tc.grid_span
        rows.append(starting)
    return rows


def _lines(block: Block) -> str:
    if isinstance(block, Paragraph):
        return block.text
    return "\n".join(cell.text for row in cells(block) for cell in row)


def runs(paragraph: Paragraph) -> list[Run]:
    inherited = _inherited(paragraph, "rPr")
    looks = []
    for run in paragraph.runs:
        levels = [run._r.rPr, *inherited]
        looks.append(
            Run(
                text=run.text,
                complex_script=_toggle(levels, "w:cs"),
                font=_font(levels, "w:ascii", "w:asciiTheme"),
                size=_size(levels, "w:sz"),
                bold=_toggle(levels, "w:b"),
                thai_font=_font(levels, "w:cs", "w:cstheme"),
                thai_size=_size(levels, "w:szCs"),
                thai_bold=_toggle(levels, "w:bCs"),
                highlight=_value(levels, "w:highlight"),
            )
        )
    return looks


def list_marker(paragraph: Paragraph) -> str | None:
    """What Word puts before a list item: "•" for any bullet, the level's text for a number
    ("%1." and the like), "" when it shows nothing. None for a paragraph outside a list."""
    numbering = None
    for properties in [paragraph._p.pPr, *_inherited(paragraph, "pPr")]:
        if properties is not None and (numbering := properties.find(qn("w:numPr"))) is not None:
            break
    if numbering is None:
        return None
    number_id = numbering.find(qn("w:numId")).get(qn("w:val"))
    level = numbering.find(qn("w:ilvl"))
    definitions = cast(DocumentPart, paragraph.part).numbering_part.element
    [number] = [n for n in definitions.findall(qn("w:num")) if n.get(qn("w:numId")) == number_id]
    abstract_id = number.find(qn("w:abstractNumId")).get(qn("w:val"))
    [abstract] = [a for a in definitions.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == abstract_id]
    ilvl = "0" if level is None else level.get(qn("w:val"))
    [definition] = [lvl for lvl in abstract.findall(qn("w:lvl")) if lvl.get(qn("w:ilvl")) == ilvl]
    text = str(definition.find(qn("w:lvlText")).get(qn("w:val")))
    return "•" if text and definition.find(qn("w:numFmt")).get(qn("w:val")) == "bullet" else text


def space_before(paragraph: Paragraph) -> float:
    """The space Word leaves above a paragraph, in points, after styles and defaults."""
    for properties in [paragraph._p.pPr, *_inherited(paragraph, "pPr")]:
        spacing = None if properties is None else properties.find(qn("w:spacing"))
        if spacing is not None and spacing.get(qn("w:before")) is not None:
            return int(spacing.get(qn("w:before"))) / 20  # twentieths of a point
    return 0.0


def _starts_page(block: Block) -> bool:
    """Word starts a new page at a paragraph set to break the page before it, and at a
    table whose first paragraph is."""
    paragraph = block if isinstance(block, Paragraph) else block.cell(0, 0).paragraphs[0]
    return bool(paragraph.paragraph_format.page_break_before)


def _inherited(paragraph: Paragraph, kind: str) -> list[BaseOxmlElement | None]:
    """Where Word looks for run ("rPr") or paragraph ("pPr") properties after the run or
    paragraph itself: the paragraph's style, the styles it is based on, then the document defaults."""
    levels: list[BaseOxmlElement | None] = []
    style = paragraph.style
    while style is not None:
        levels.append(style.element.find(qn(f"w:{kind}")))
        style = style.base_style
    defaults = cast(DocumentPart, paragraph.part).styles.element.find(qn("w:docDefaults"))
    levels.append(None if defaults is None else defaults.find(f"{qn(f'w:{kind}Default')}/{qn(f'w:{kind}')}"))
    return levels


def _first(levels: list[BaseOxmlElement | None], tag: str) -> BaseOxmlElement | None:
    for level in levels:
        if level is not None and (element := level.find(qn(tag))) is not None:
            return cast(BaseOxmlElement, element)
    return None


def _toggle(levels: list[BaseOxmlElement | None], tag: str) -> bool:
    element = _first(levels, tag)
    return element is not None and element.get(qn("w:val")) not in ("0", "false")


def _value(levels: list[BaseOxmlElement | None], tag: str) -> str | None:
    element = _first(levels, tag)
    return None if element is None else str(element.get(qn("w:val")))


def _size(levels: list[BaseOxmlElement | None], tag: str) -> float:
    element = _first(levels, tag)
    return 0.0 if element is None else int(element.get(qn("w:val"), "0")) / 2  # half-points


def _font(levels: list[BaseOxmlElement | None], name: str, theme: str) -> str:
    """The font named for a slot. A theme font at the same level wins over the name."""
    for level in levels:
        fonts = None if level is None else level.find(qn("w:rFonts"))
        if fonts is None:
            continue
        if fonts.get(qn(theme)):
            return f"theme {fonts.get(qn(theme))}"
        if fonts.get(qn(name)):
            return str(fonts.get(qn(name)))
    return ""
