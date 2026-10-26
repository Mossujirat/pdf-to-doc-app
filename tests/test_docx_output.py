"""DOCX output: an editable Word file with the same text and structure as the Markdown output."""

import io
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx_reader import Block, Cell, cells, list_marker, page_texts, read_docx, runs, space_before
from pdf_builder import Cell as GridCell
from pdf_builder import Grid, Line, Styled, pdf

from pdf_to_doc import convert
from pdf_to_doc.cli import main


def docx_of(*pages: list[Line]) -> bytes:
    return convert(pdf(*pages), to="docx").output


def paragraphs(*pages: list[Line]) -> list[list[Paragraph]]:
    """The paragraphs of each page, for pages that hold no table."""
    blocks = read_docx(docx_of(*pages)).pages
    assert all(isinstance(block, Paragraph) for page in blocks for block in page)
    return [[block for block in page if isinstance(block, Paragraph)] for page in blocks]


def shown(block: Block) -> str | list[list[Cell]]:
    """A paragraph's text, or a table's cells."""
    return block.text if isinstance(block, Paragraph) else cells(block)


def test_a_thai_text_page_comes_out_as_its_text() -> None:
    assert page_texts(docx_of(["ภาษาไทยอ่านได้ถูกต้อง"])) == ["ภาษาไทยอ่านได้ถูกต้อง"]


def test_each_page_of_the_original_starts_a_new_page_after_the_summary() -> None:
    output = read_docx(docx_of(["หน้าแรก"], ["หน้าที่สอง"], ["หน้าที่สาม"]))

    assert output.summary.text.startswith("สรุปการแปลง 3 หน้า")
    assert [[shown(block) for block in page] for page in output.pages] == [
        ["หน้าแรก"],
        ["หน้าที่สอง"],
        ["หน้าที่สาม"],
    ]


def test_a_page_that_starts_with_a_table_starts_a_new_page() -> None:
    grid = Grid(rows=(("งวดที่", "จำนวนเงิน"), ("๑", "50,000")), widths=(150, 150))

    pages = read_docx(docx_of(["หน้าแรก"], [grid, "หน้าที่สอง"])).pages

    assert [[shown(block) for block in page] for page in pages] == [
        ["หน้าแรก"],
        [[[Cell("งวดที่"), Cell("จำนวนเงิน")], [Cell("๑"), Cell("50,000")]], "หน้าที่สอง"],
    ]


def test_no_paragraph_is_left_empty_to_carry_a_page_break() -> None:
    # Found in real Word: when the page before was full, such a paragraph fell onto
    # a page of its own and left that page blank.
    word = docx.Document(io.BytesIO(docx_of(["หน้าแรก"], ["หน้าที่สอง"], ["หน้าที่สาม"])))

    assert [paragraph.text for paragraph in word.paragraphs if not paragraph.text] == []


def test_pages_are_a4() -> None:
    [section] = docx.Document(io.BytesIO(docx_of(["หน้าแรก"]))).sections

    assert section.page_width is not None and section.page_height is not None
    assert (round(section.page_width.mm), round(section.page_height.mm)) == (210, 297)


def test_the_summary_opens_with_its_label_in_bold() -> None:
    summary = read_docx(docx_of(["หน้าแรก"])).summary

    [label, *rest] = runs(summary)
    assert (label.text, label.thai_bold) == ("สรุปการแปลง", True)
    assert "".join(run.text for run in rest).strip() == "1 หน้า: อ่านจากข้อความในไฟล์ 1 หน้า, อ่านด้วย OCR 0 หน้า"
    assert not any(run.thai_bold for run in rest)


# Fonts


def test_thai_text_is_set_in_th_sarabun_new_16_pt_in_the_complex_script_slot() -> None:
    [[paragraph]] = paragraphs(["ภาษาไทยอ่านได้ถูกต้อง"])

    [run] = runs(paragraph)
    assert (run.text, run.complex_script, run.thai_font, run.thai_size) == (
        "ภาษาไทยอ่านได้ถูกต้อง",
        True,
        "TH Sarabun New",
        16,
    )


def test_english_words_get_runs_of_their_own_that_are_not_complex_script() -> None:
    [[paragraph]] = paragraphs(["พัฒนาระบบเศรษฐกิจชีวภาพ (Bio-Circular-Green Economy) ปี 2569"])

    assert [(run.text, run.complex_script) for run in runs(paragraph)] == [
        ("พัฒนาระบบเศรษฐกิจชีวภาพ (", True),
        ("Bio-Circular-Green Economy", False),
        (") ปี 2569", True),
    ]


def test_english_text_is_set_in_th_sarabun_new_at_the_same_size_as_thai() -> None:
    [[paragraph]] = paragraphs(["ภาษาไทย English ภาษาไทย"])

    assert {(run.font, run.size, run.thai_font, run.thai_size) for run in runs(paragraph)} == {
        ("TH Sarabun New", 16, "TH Sarabun New", 16)
    }


# Structure


def test_headings_use_words_heading_styles_at_their_levels_in_bold_th_sarabun_new() -> None:
    [page] = paragraphs(
        [
            Styled("สัญญาให้ทุนโครงการ", size=24),
            Styled("ข้อกำหนดทั่วไป", size=20),
            "ผู้ให้ทุนตกลงให้ทุนแก่ผู้รับทุน",
        ]
    )

    assert [(paragraph.style.name if paragraph.style else "", paragraph.text) for paragraph in page] == [
        ("Heading 1", "สัญญาให้ทุนโครงการ"),
        ("Heading 2", "ข้อกำหนดทั่วไป"),
        ("Normal", "ผู้ให้ทุนตกลงให้ทุนแก่ผู้รับทุน"),
    ]
    looks = [
        {(run.thai_font, run.thai_size, run.thai_bold, run.font, run.size, run.bold) for run in runs(heading)}
        for heading in page[:2]
    ]
    assert looks == [
        {("TH Sarabun New", 20, True, "TH Sarabun New", 20, True)},
        {("TH Sarabun New", 18, True, "TH Sarabun New", 18, True)},
    ]


def test_bulleted_items_get_words_bullets_and_numbered_items_keep_only_their_own_numbers() -> None:
    [page] = paragraphs(
        [
            "เอกสารแนบท้ายสัญญามีดังนี้",
            "1. เงื่อนไขของสัญญา",
            "(ก) ขอบเขตของการดำเนินงาน",
            "",
            "• รายงานความก้าวหน้า",
            "• รายงานฉบับสมบูรณ์",
        ]
    )

    assert [(list_marker(paragraph), paragraph.text) for paragraph in page] == [
        (None, "เอกสารแนบท้ายสัญญามีดังนี้"),
        ("", "1. เงื่อนไขของสัญญา"),
        ("", "(ก) ขอบเขตของการดำเนินงาน"),
        ("•", "รายงานความก้าวหน้า"),
        ("•", "รายงานฉบับสมบูรณ์"),
    ]


def test_a_ruled_table_becomes_a_word_table_between_the_text_around_it() -> None:
    grid = Grid(rows=(("งวดที่", "จำนวนเงิน"), ("๑", "50,000"), ("๒", "75,000")), widths=(150, 150))

    [page] = read_docx(docx_of(["ผู้ให้ทุนจะจ่ายเงินเป็นงวด", grid, "ตามเงื่อนไขในสัญญา"])).pages

    assert [shown(block) for block in page] == [
        "ผู้ให้ทุนจะจ่ายเงินเป็นงวด",
        [[Cell("งวดที่"), Cell("จำนวนเงิน")], [Cell("๑"), Cell("50,000")], [Cell("๒"), Cell("75,000")]],
        "ตามเงื่อนไขในสัญญา",
    ]


def test_merged_cells_are_merged_in_the_word_table() -> None:
    grid = Grid(
        rows=(
            (GridCell("รายการ", rowspan=2), GridCell("งบประมาณ", colspan=2)),
            ("ปีที่ ๑", "ปีที่ ๒"),
            ("ค่าตอบแทน", "100,000", "120,000"),
        ),
        widths=(150, 100, 100),
    )

    [[table]] = read_docx(docx_of([grid])).pages

    assert isinstance(table, Table)
    assert cells(table) == [
        [Cell("รายการ", rowspan=2), Cell("งบประมาณ", colspan=2)],
        [Cell("ปีที่ ๑"), Cell("ปีที่ ๒")],
        [Cell("ค่าตอบแทน"), Cell("100,000"), Cell("120,000")],
    ]


def test_sub_rows_are_merged_and_items_keep_their_line_breaks_in_the_word_table() -> None:
    grid = Grid(
        rows=(
            ("หมวด", "รายการ", "จำนวนเงิน"),
            ("ค่าดำเนินงาน", "ค่าวัสดุสำนักงาน\nค่าถ่ายเอกสาร", "12,000\n3,500"),
            ("ค่าครุภัณฑ์", "เครื่องคอมพิวเตอร์\nเครื่องพิมพ์", "53,900"),
        ),
        widths=(110, 210, 100),
    )

    [[table]] = read_docx(docx_of([grid])).pages

    assert isinstance(table, Table)
    assert cells(table) == [
        [Cell("หมวด"), Cell("รายการ"), Cell("จำนวนเงิน")],
        [Cell("ค่าดำเนินงาน", rowspan=2), Cell("ค่าวัสดุสำนักงาน"), Cell("12,000")],
        [Cell("ค่าถ่ายเอกสาร"), Cell("3,500")],
        [Cell("ค่าครุภัณฑ์"), Cell("เครื่องคอมพิวเตอร์\nเครื่องพิมพ์"), Cell("53,900")],
    ]


def test_text_after_a_table_keeps_clear_of_its_bottom_line() -> None:
    # Found in real Word: the next line sat on the table's bottom border.
    grid = Grid(rows=(("งวดที่", "จำนวนเงิน"), ("๑", "50,000")), widths=(150, 150))

    [[_, paragraph]] = read_docx(docx_of([grid, "ตามเงื่อนไขในสัญญา"])).pages

    assert isinstance(paragraph, Paragraph)
    assert space_before(paragraph) >= 10


def test_two_tables_in_a_row_stay_two_tables() -> None:
    # Word joins tables that touch into one.
    first = Grid(rows=(("งวดที่", "จำนวนเงิน"), ("๑", "50,000")), widths=(150, 150))
    second = Grid(rows=(("รายการ", "หมายเหตุ"), ("ค่าวัสดุ", "เบิกได้")), widths=(150, 150))

    output = docx_of([first, "", second])

    [page] = read_docx(output).pages
    assert [shown(block) for block in page] == [
        [[Cell("งวดที่"), Cell("จำนวนเงิน")], [Cell("๑"), Cell("50,000")]],
        [[Cell("รายการ"), Cell("หมายเหตุ")], [Cell("ค่าวัสดุ"), Cell("เบิกได้")]],
    ]
    body = list(docx.Document(io.BytesIO(output)).element.body)
    assert not any(
        before.tag == after.tag == qn("w:tbl") for before, after in zip(body, body[1:])
    ), "nothing between the tables"


# The command line


def test_the_command_writes_a_docx_beside_the_source_when_asked_for_one(tmp_path: Path) -> None:
    source = tmp_path / "contract.pdf"
    source.write_bytes(pdf(["สัญญาให้ทุนโครงการ"]))

    assert main([str(source), "--to", "docx"]) == 0
    assert page_texts((tmp_path / "contract.docx").read_bytes()) == ["สัญญาให้ทุนโครงการ"]


def test_an_output_file_named_docx_gets_a_docx_without_asking(tmp_path: Path) -> None:
    source = tmp_path / "contract.pdf"
    source.write_bytes(pdf(["สัญญาให้ทุนโครงการ"]))

    assert main([str(source), "-o", str(tmp_path / "แก้ไข.DOCX")]) == 0
    assert page_texts((tmp_path / "แก้ไข.DOCX").read_bytes()) == ["สัญญาให้ทุนโครงการ"]
