"""Ruled tables on OCR pages come out as tables: the rules are found in the page image,
and the text the engine reads is placed in the cells they make."""

from collections.abc import Sequence
from typing import Literal

import pytest
from docx.table import Table
from docx_reader import Cell as WordCell
from docx_reader import cells, read_docx, runs
from fake_ocr import Askew, FakeEngine, Sideways, Boxed, Unsure, marks_over
from fake_ocr import Line as ReadLine
from generate import (
    BUDGET,
    BUDGET_REFERENCE,
    COARSE_TABLE,
    COARSE_TABLE_REFERENCE,
    DATES,
    DATES_REFERENCE,
    KAR_AS_KAM,
    PAYMENT_SCHEDULE,
    PAYMENT_SCHEDULE_REFERENCE,
    PRICE_SCHEDULE,
    PRICE_SCHEDULE_REFERENCE,
    scanned,
)
from markdown_reader import page_texts
from pdf_builder import (
    FONT_SIZE,
    LEFT,
    LINE_HEIGHT,
    TOP,
    Cell,
    Drawn,
    Grid,
    Line,
    Picture,
    Scan,
    Styled,
    laid_out_grid,
    pdf,
    remapped,
)
from test_ocr_pages import THAI_LEFT, THAI_RIGHT, two_columns

from pdf_to_doc import convert


def markdown_pages(scanned_lines: Sequence[Line], read: Sequence[ReadLine]) -> list[str]:
    """The Markdown of a scan of the lines, scanned as the corpus's scans are, which the engine reads as `read`."""
    return converted(pdf(scanned(scanned_lines, seed=1)), FakeEngine(read))


def converted(source: bytes, engine: FakeEngine) -> list[str]:
    return page_texts(convert(source, ocr=engine).output.decode("utf-8"))


def schedule(task: str) -> Grid:
    return Grid(rows=(("งวดที่", "งาน"), ("๑", task)), widths=(80, 300))


def test_a_ruled_table_on_a_scan_becomes_a_table_between_the_text_around_it() -> None:
    grid = Grid(rows=(("งวดที่", "จำนวนเงิน"), ("๑", "50,000"), ("๒", "75,000")), widths=(150, 150))
    lines: list[str | Grid] = ["ผู้ให้ทุนจะจ่ายเงินเป็นงวด", grid, "ตามเงื่อนไขในสัญญา"]

    pages = markdown_pages(lines, lines)

    assert pages == [
        "ผู้ให้ทุนจะจ่ายเงินเป็นงวด\n\n"
        "| งวดที่ | จำนวนเงิน |\n"
        "| --- | --- |\n"
        "| ๑ | 50,000 |\n"
        "| ๒ | 75,000 |\n\n"
        "ตามเงื่อนไขในสัญญา"
    ]


def test_a_scanned_table_keeps_both_lines_of_a_wrapped_cell_and_its_merged_cells() -> None:
    lines: list[str | Grid] = ["ผู้ให้ทุนจะจ่ายเงินเป็นสามงวด", PAYMENT_SCHEDULE, "งบประมาณแยกตามปี", BUDGET]

    pages = markdown_pages(lines, lines)

    assert pages == [
        f"ผู้ให้ทุนจะจ่ายเงินเป็นสามงวด\n\n{PAYMENT_SCHEDULE_REFERENCE}\n\nงบประมาณแยกตามปี\n\n{BUDGET_REFERENCE}"
    ]


def test_several_printed_lines_in_a_scanned_cell_are_read_as_on_a_text_page() -> None:
    # Items stacked beside their amounts are sub-rows; items beside one total keep their breaks;
    # a wrapped description is one sentence.
    lines: list[str | Grid] = [COARSE_TABLE]

    pages = markdown_pages(lines, lines)

    assert pages == [COARSE_TABLE_REFERENCE]


def test_a_scanned_table_is_a_word_table_with_its_merged_cells_and_breaks() -> None:
    lines: list[str | Grid] = [COARSE_TABLE]

    result = convert(pdf(scanned(lines, seed=1)), to="docx", ocr=FakeEngine(lines))

    [[table]] = read_docx(result.output).pages
    assert isinstance(table, Table)
    assert cells(table) == [
        [WordCell("หมวดรายจ่าย"), WordCell("รายการ"), WordCell("จำนวนเงิน")],
        [WordCell("ค่าดำเนินงาน", rowspan=3), WordCell("ค่าวัสดุสำนักงาน"), WordCell("12,000")],
        [WordCell("ค่าถ่ายเอกสาร"), WordCell("3,500")],
        [WordCell("ค่าไปรษณีย์"), WordCell("1,200")],
        [WordCell("ค่าครุภัณฑ์"), WordCell("เครื่องคอมพิวเตอร์\nเครื่องพิมพ์"), WordCell("53,900")],
        [WordCell("ค่าเดินทาง"), WordCell("ค่าเช่ารถและค่าน้ำมันเชื้อเพลิงสำหรับการลงพื้นที่"), WordCell("15,000")],
    ]


def test_a_scanned_table_of_two_rows_keeps_each_date_under_its_heading() -> None:
    lines: list[str | Grid] = ["ระยะเวลาของโครงการ", DATES]

    pages = markdown_pages(lines, lines)

    assert pages == [f"ระยะเวลาของโครงการ\n\n{DATES_REFERENCE}"]


def test_a_double_rule_above_a_total_is_read_as_one_rule() -> None:
    # A row of empty cells 3 points tall draws the second rule.
    grid = Grid(
        rows=(("รายการ", "จำนวนเงิน"), ("ค่าวัสดุ", "45,000"), ("ค่าเดินทาง", "12,000"), ("", ""), ("รวม", "57,000")),
        widths=(170, 110),
        heights=((3, 3.0),),
    )

    pages = markdown_pages([grid], [grid])

    assert pages == [
        "| รายการ | จำนวนเงิน |\n| --- | --- |\n| ค่าวัสดุ | 45,000 |\n| ค่าเดินทาง | 12,000 |\n| รวม | 57,000 |"
    ]


@pytest.mark.parametrize("angle", [-3.0, 3.0])
def test_a_table_scanned_askew_comes_out_as_a_table(angle: float) -> None:
    # The engine straightens its boxes by the page's slant, and the rules are straightened to match.
    lines: list[str | Grid] = ["งบประมาณแยกตามปี", BUDGET]
    source = pdf(Scan(lines, dpi=200, angle=angle, noise=6, seed=1))

    pages = converted(source, FakeEngine(Askew(lines, angle)))

    assert pages == [f"งบประมาณแยกตามปี\n\n{BUDGET_REFERENCE}"]


@pytest.mark.parametrize("turn", [90, 270])
def test_a_table_printed_sideways_comes_out_as_a_table_once_the_page_is_turned(turn: Literal[90, 270]) -> None:
    pages = converted(pdf(scanned(PRICE_SCHEDULE, seed=1, sideways=turn)), FakeEngine(Sideways(PRICE_SCHEDULE, turn)))

    assert pages == [PRICE_SCHEDULE_REFERENCE]


# OCR pages that are not scans: rendered from the PDF, whose rules are in the image too

@pytest.mark.parametrize("keep_images", [False, True], ids=["images read", "images kept"])
@pytest.mark.parametrize("drawn", [False, True], ids=["text layer that cannot be trusted", "text drawn as shapes"])
def test_a_ruled_table_on_an_ocr_page_that_is_not_a_scan_becomes_a_table(drawn: bool, keep_images: bool) -> None:
    # Kept images keep scans only: these pages are still read.
    lines: list[str | Grid] = ["ผู้ให้ทุนจะจ่ายเงินเป็นสามงวด", PAYMENT_SCHEDULE, "งบประมาณแยกตามปี", BUDGET]
    source = pdf(Drawn(lines)) if drawn else remapped(pdf(lines), KAR_AS_KAM)

    result = convert(source, ocr=FakeEngine(lines), keep_images=keep_images)

    assert [page.kind for page in result.summary.pages] == ["ocr"]
    assert page_texts(result.output.decode("utf-8")) == [
        f"ผู้ให้ทุนจะจ่ายเงินเป็นสามงวด\n\n{PAYMENT_SCHEDULE_REFERENCE}\n\nงบประมาณแยกตามปี\n\n{BUDGET_REFERENCE}"
    ]


def test_a_scanned_table_kept_as_an_image_is_not_read() -> None:
    lines: list[str | Grid] = ["ผู้ให้ทุนจะจ่ายเงินเป็นสามงวด", PAYMENT_SCHEDULE]
    engine = FakeEngine(lines)

    result = convert(pdf(scanned(lines, seed=1)), ocr=engine, keep_images=True)

    assert page_texts(result.output.decode("utf-8")) == ["[หน้าสแกน]"]
    assert engine.images == []


def test_a_note_typed_over_a_scanned_table_lands_in_the_cell_it_sits_in() -> None:
    grid = Grid(rows=(("งวดที่", "จำนวนเงิน", "หมายเหตุ"), ("๑", "50,000", ""), ("๒", "75,000", "")), widths=(80, 120, 200))
    # In the first row's empty note cell; the engine reads the note too, with a letter lost.
    note, misread_note = Styled("แก้ไขตามบันทึกแล้ว", x=276, y=TOP + 53), Styled("แกไขตามบันทึกแล้ว", x=276, y=TOP + 53)
    lines: list[str | Grid] = ["ผู้ให้ทุนจะจ่ายเงินเป็นงวด", grid]

    pages = converted(pdf(Scan(lines, overlay=[note], dpi=200, angle=0.4, noise=6, seed=1)), FakeEngine([*lines, misread_note]))

    assert pages == [
        "ผู้ให้ทุนจะจ่ายเงินเป็นงวด\n\n"
        "| งวดที่ | จำนวนเงิน | หมายเหตุ |\n"
        "| --- | --- | --- |\n"
        "| ๑ | 50,000 | แก้ไขตามบันทึกแล้ว |\n"
        "| ๒ | 75,000 |  |"
    ]


def test_a_table_right_under_two_columns_does_not_stop_them_being_read_one_after_the_other() -> None:
    # The table starts a line under the columns and reaches across their gutter: taken for text,
    # its cells would close the gutter, and the page would be read across.
    table = Grid(
        rows=(("หมวด", "รายละเอียดค่าใช้จ่ายของโครงการ"), ("ค่าตอบแทน", "ค่าตอบแทนผู้วิจัยและผู้ช่วยวิจัย")),
        widths=(150, 306),
        y=TOP + 5 * LINE_HEIGHT,
    )
    lines: list[str | Styled | Grid] = [*two_columns(THAI_LEFT, THAI_RIGHT, y=TOP), table]

    [page] = markdown_pages(lines, lines)

    assert page == "\n\n".join(
        [
            "".join(THAI_LEFT),
            "".join(THAI_RIGHT),
            "| หมวด | รายละเอียดค่าใช้จ่ายของโครงการ |\n| --- | --- |\n| ค่าตอบแทน | ค่าตอบแทนผู้วิจัยและผู้ช่วยวิจัย |",
        ]
    )


def test_a_table_in_a_column_stays_whole_and_the_columns_are_read_one_after_the_other() -> None:
    # The left column holds a table; the right column runs on beside it. The table's own short
    # cells, taken for text, would part the left column in three, and the page would be read across.
    above, below = THAI_LEFT[:2], THAI_LEFT[2:]
    right = [*THAI_RIGHT, *THAI_LEFT[:4]]
    table = Grid(rows=(("ปี", "งบประมาณ"), ("2568", "50,000")), widths=(90, 110), y=TOP + 3 * LINE_HEIGHT)
    lines: list[str | Styled | Grid] = [
        *(Styled(text, size=12, y=TOP + number * LINE_HEIGHT) for number, text in enumerate(above)),
        table,
        *(Styled(text, size=12, y=TOP + (6 + number) * LINE_HEIGHT) for number, text in enumerate(below)),
        *(Styled(text, size=12, x=300, y=TOP + number * LINE_HEIGHT) for number, text in enumerate(right)),
    ]

    [page] = markdown_pages(lines, lines)

    assert page == "\n\n".join(["".join(above), "| ปี | งบประมาณ |\n| --- | --- |\n| 2568 | 50,000 |", "".join(below), "".join(right)])


def test_a_table_in_the_right_column_is_read_in_that_column_where_it_sits() -> None:
    # The left column's second paragraph starts below the table's top, but beside it.
    table = Grid(rows=(("ปี", "งบประมาณ"), ("2568", "50,000")), widths=(90, 110), x=300, y=TOP + 3 * LINE_HEIGHT)
    left = [THAI_LEFT[:2], THAI_LEFT[2:]]
    above, below = THAI_RIGHT[:2], [*THAI_RIGHT[2:], *THAI_LEFT[:3]]
    lines: list[str | Styled | Grid] = [
        *(Styled(text, size=12, y=TOP + number * LINE_HEIGHT) for number, text in enumerate(left[0])),
        *(Styled(text, size=12, y=TOP + (3 + number) * LINE_HEIGHT) for number, text in enumerate(left[1])),
        *(Styled(text, size=12, x=300, y=TOP + number * LINE_HEIGHT) for number, text in enumerate(above)),
        table,
        *(Styled(text, size=12, x=300, y=TOP + (6 + number) * LINE_HEIGHT) for number, text in enumerate(below)),
    ]

    [page] = markdown_pages(lines, lines)

    assert page == "\n\n".join(
        ["".join(left[0]), "".join(left[1]), "".join(above), "| ปี | งบประมาณ |\n| --- | --- |\n| 2568 | 50,000 |", "".join(below)]
    )


# Rules that do not make a table


def test_a_frame_around_a_paragraph_on_a_scan_is_boxed_text_not_a_table() -> None:
    notice = Grid(rows=((Cell("ผู้รับทุนต้องส่งรายงานความก้าวหน้า\nทุกหกเดือนนับแต่วันลงนาม\nในสัญญาฉบับนี้"),),), widths=(400,))

    pages = markdown_pages([notice], [notice])

    assert pages == ["ผู้รับทุนต้องส่งรายงานความก้าวหน้าทุกหกเดือนนับแต่วันลงนามในสัญญาฉบับนี้"]


def test_rules_that_meet_but_close_off_no_space_make_no_table() -> None:
    # A rule across the page and one down its left side, meeting at the corner, drawn as cells a
    # point thin: there is no space between them for a cell.
    corner = Grid(rows=(("", ""), ("", Cell("", boxed=False))), widths=(1, 300), heights=((0, 1.0), (1, 120.0)))
    lines: list[str | Grid] = [corner, "ผู้รับทุนต้องส่งรายงานทุกหกเดือน"]

    pages = markdown_pages(lines, lines)

    assert pages == ["ผู้รับทุนต้องส่งรายงานทุกหกเดือน"]


# Tables whose text cannot be placed in cells with certainty: incomplete tables

def boxed_header_table(baseline: float, *, misread: bool = False) -> list[Styled | Grid]:
    """A table whose header alone is boxed, its first line on the baseline, and its rows set out
    under it in the header's columns, ruled at 72, 252, 372 and 492 points. `misread` reads
    "เดือน" as "เคือน", which no word list holds."""
    month = "เคือน" if misread else "เดือน"
    rows = (("ระบบติดตามงาน", "1,200,000", f"12 {month}"), ("ระบบจัดเก็บเอกสาร", "800,000", f"6 {month}"))
    return [
        Grid(rows=(("ชื่อโครงการ", "งบประมาณ", "ระยะเวลา"),), widths=(180, 120, 120), y=baseline),
        *(
            Styled(text, x=x, y=baseline + 38 + number * LINE_HEIGHT)
            for number, row in enumerate(rows)
            for x, text in zip((76, 256, 376), row)
        ),
    ]


BOXED_HEADER_ROWS = ["ชื่อโครงการ งบประมาณ ระยะเวลา", "ระบบติดตามงาน 1,200,000 12 เดือน", "ระบบจัดเก็บเอกสาร 800,000 6 เดือน"]
UNDER_IT = TOP + 38 + 2 * LINE_HEIGHT  # the baseline under such a table on a page's first line


def test_a_table_whose_header_alone_is_boxed_is_an_incomplete_table_of_its_rows() -> None:
    # The paragraph under the rows runs across the header's first column rule: it is no row of the
    # table. Taken for rows, its two lines would be parted.
    after = [Styled("ทั้งนี้ระยะเวลานับแต่วันลงนาม", y=UNDER_IT + LINE_HEIGHT), Styled("ในสัญญาฉบับนี้")]
    lines: list[str | Styled | Grid] = ["โครงการที่เสนอขอรับทุน", *boxed_header_table(TOP + LINE_HEIGHT), *after]

    pages = markdown_pages(lines, lines)

    assert pages == [
        "\n\n".join(["โครงการที่เสนอขอรับทุน", "[ตารางอาจไม่สมบูรณ์]", *BOXED_HEADER_ROWS, "ทั้งนี้ระยะเวลานับแต่วันลงนามในสัญญาฉบับนี้"])
    ]


def test_an_incomplete_table_ends_where_a_gap_wider_than_a_line_parts_it_from_what_follows() -> None:
    # Two short lines of a signature block, two lines under the table, within its first column.
    after = [Styled("ลงชื่อผู้เสนอโครงการ", y=UNDER_IT + 3 * LINE_HEIGHT), Styled("หัวหน้าโครงการ")]
    lines: list[Styled | Grid] = [*boxed_header_table(TOP), *after]

    pages = markdown_pages(lines, lines)

    assert pages == ["\n\n".join(["[ตารางอาจไม่สมบูรณ์]", *BOXED_HEADER_ROWS, "ลงชื่อผู้เสนอโครงการหัวหน้าโครงการ"])]


def test_the_wrapped_line_of_a_cell_under_a_boxed_header_stays_in_the_incomplete_table() -> None:
    # Only the header is boxed; the second row's description wraps onto a line of its own column,
    # on the baseline at TOP + 81. Read outside the table, it would run on into the sentence under it.
    def unboxed(*texts: str) -> tuple[Cell, ...]:
        return tuple(Cell(text, boxed=False) for text in texts)

    grid = Grid(
        rows=(
            ("งวดที่", "รายละเอียดงานที่ส่งมอบ", "จำนวนเงิน"),
            unboxed("๑", "ส่งแผนการดำเนินงาน", "150,000"),
            unboxed("๒", "ส่งรายงานความก้าวหน้า\nพร้อมผลการทดลอง", "250,000"),
        ),
        widths=(60, 230, 130),
    )
    lines: list[Styled | Grid] = [grid, Styled("ทั้งนี้ไม่รวมภาษีมูลค่าเพิ่มตามที่กฎหมายกำหนด", y=TOP + 81 + LINE_HEIGHT)]

    pages = markdown_pages(lines, lines)

    assert pages == [
        "[ตารางอาจไม่สมบูรณ์]\n\nงวดที่ รายละเอียดงานที่ส่งมอบ จำนวนเงิน\n\n๑ ส่งแผนการดำเนินงาน 150,000\n\n"
        "๒ ส่งรายงานความก้าวหน้า 250,000\n\nพร้อมผลการทดลอง\n\nทั้งนี้ไม่รวมภาษีมูลค่าเพิ่มตามที่กฎหมายกำหนด"
    ]


def test_an_incomplete_table_in_a_column_takes_none_of_the_column_beside_it() -> None:
    # The left column holds a table whose header alone is boxed; the right column runs on beside
    # its rows, level with them.
    above, below = THAI_LEFT[:2], THAI_LEFT[2:]
    right = [*THAI_RIGHT, *THAI_LEFT[:4]]
    header = Grid(rows=(("ปี", "งบประมาณ"),), widths=(90, 110), y=TOP + 3 * LINE_HEIGHT)
    rows = [("2568", "50,000"), ("2569", "60,000")]
    lines: list[str | Styled | Grid] = [
        *(Styled(text, size=12, y=TOP + number * LINE_HEIGHT) for number, text in enumerate(above)),
        header,
        *(
            Styled(text, size=12, x=x, y=TOP + 3 * LINE_HEIGHT + 38 + number * LINE_HEIGHT)
            for number, row in enumerate(rows)
            for x, text in zip((LEFT + 4, LEFT + 94), row)
        ),
        *(Styled(text, size=12, y=TOP + (7 + number) * LINE_HEIGHT) for number, text in enumerate(below)),
        *(Styled(text, size=12, x=300, y=TOP + number * LINE_HEIGHT) for number, text in enumerate(right)),
    ]

    [page] = markdown_pages(lines, lines)

    assert page == "\n\n".join(
        ["".join(above), "[ตารางอาจไม่สมบูรณ์]", "ปี งบประมาณ", "2568 50,000", "2569 60,000", "".join(below), "".join(right)]
    )


def test_misread_words_in_an_incomplete_table_are_marked_and_counted() -> None:
    result = convert(pdf(scanned(boxed_header_table(TOP), seed=1)), ocr=FakeEngine(boxed_header_table(TOP, misread=True)))

    assert page_texts(result.output.decode("utf-8")) == [
        "[ตารางอาจไม่สมบูรณ์]\n\nชื่อโครงการ งบประมาณ ระยะเวลา\n\n"
        "ระบบติดตามงาน 1,200,000 12 <mark>เคือน</mark>\n\nระบบจัดเก็บเอกสาร 800,000 6 <mark>เคือน</mark>"
    ]
    assert [page.uncertain for page in result.summary.pages] == [2]


def test_an_incomplete_table_in_word_is_its_label_then_a_paragraph_per_row_highlighted_where_uncertain() -> None:
    result = convert(
        pdf(scanned(boxed_header_table(TOP), seed=1)), to="docx", ocr=FakeEngine(boxed_header_table(TOP, misread=True))
    )

    [page] = read_docx(result.output).pages
    assert [block.text for block in page if not isinstance(block, Table)] == [
        "[ตารางอาจไม่สมบูรณ์]",
        BOXED_HEADER_ROWS[0],
        *(row.replace("เดือน", "เคือน") for row in BOXED_HEADER_ROWS[1:]),
    ]
    assert [[run.text for run in runs(block) if run.highlight] for block in page if not isinstance(block, Table)] == [
        [],
        [],
        ["เคือน"],
        ["เคือน"],
    ]


def test_the_summary_counts_tables_on_ocr_pages_and_incomplete_ones_with_their_pages() -> None:
    # A text page with a table, a scan with neither, and a scan with a table and an incomplete one.
    text_page: list[str | Styled | Grid] = [Grid(rows=(("ปี", "งบ"), ("2568", "100")), widths=(100, 100))]
    plain: list[str | Styled | Grid] = ["ผู้รับทุนต้องส่งรายงานตามกำหนด"]
    tables: list[str | Styled | Grid] = [schedule("ส่งรายงานทุกหกเดือน"), *boxed_header_table(TOP + 5 * LINE_HEIGHT)]

    result = convert(pdf(text_page, scanned(plain, seed=1), scanned(tables, seed=2)), ocr=FakeEngine(plain, tables))

    assert [(page.tables, page.incomplete_tables) for page in result.summary.pages] == [(0, 0), (0, 0), (1, 1)]
    assert result.summary.headline().endswith(
        " ตารางในหน้าที่อ่านด้วย OCR: เป็นตาราง 1 ตาราง, อาจไม่สมบูรณ์ 1 ตาราง ในหน้า 3 ควรตรวจโครงสร้างกับต้นฉบับ"
    )


@pytest.mark.parametrize(
    ("lines", "said"),
    [
        (["ผู้รับทุนต้องส่งรายงานตามกำหนด"], ""),
        ([schedule("ส่งรายงานทุกหกเดือน")], " ตารางในหน้าที่อ่านด้วย OCR: เป็นตาราง 1 ตาราง"),
        (boxed_header_table(TOP), " ตารางในหน้าที่อ่านด้วย OCR: อาจไม่สมบูรณ์ 1 ตาราง ในหน้า 2 ควรตรวจโครงสร้างกับต้นฉบับ"),
    ],
    ids=["a text page's table alone", "a table", "an incomplete table"],
)
def test_the_headline_says_only_the_kinds_of_ocr_table_there_are(lines: list[str | Styled | Grid], said: str) -> None:
    # After a text page with a table, which is not counted.
    text_page: list[str | Styled | Grid] = [Grid(rows=(("ปี", "งบ"), ("2568", "100")), widths=(100, 100))]

    result = convert(pdf(text_page, scanned(lines, seed=1)), ocr=FakeEngine(lines))

    headline = result.summary.headline()
    assert headline.endswith(said) and headline.count("ตาราง") == said.count("ตาราง")


@pytest.mark.parametrize(
    "under",
    ["ผู้รับทุนลงนามรับรองว่าข้อมูลข้างต้นเป็นความจริงทุกประการ", "กรอกให้ครบทุกช่อง"],
    ids=["a sentence across its columns", "a note in one of them"],
)
def test_a_boxed_row_with_no_row_set_out_under_it_is_read_as_lines(under: str) -> None:
    # Boxes of a form, with text under them that is set out in no more than one of its columns.
    header, *_ = boxed_header_table(TOP)
    lines: list[str | Styled | Grid] = [header, Styled(under, y=TOP + 38)]

    pages = markdown_pages(lines, lines)

    assert pages == [f"{BOXED_HEADER_ROWS[0]}\n\n{under}"]


def test_a_box_that_runs_through_a_rule_makes_its_table_incomplete() -> None:
    # The engine boxes the second row's two cells as one line, through the rule between them.
    scanned_grid = Grid(rows=(("งวดที่", "งาน"), ("๑", "ส่งรายงานทุกหกเดือน")), widths=(80, 300))
    read_grid = Grid(rows=(("งวดที่", "งาน"), (Cell("๑ ส่งรายงานทุกหกเดือน", colspan=2),)), widths=(80, 300))

    pages = markdown_pages([scanned_grid], [read_grid])

    assert pages == ["[ตารางอาจไม่สมบูรณ์]\n\nงวดที่ งาน\n\n๑ ส่งรายงานทุกหกเดือน"]


def test_a_box_that_runs_out_through_the_frame_of_a_table_makes_it_incomplete() -> None:
    # The engine boxes the last cell's text past the table's right rule, 380 points from its left,
    # as it would with a note printed beside the table: what lies outside cannot be told apart.
    grid = Grid(rows=(("งวดที่", "งาน"), ("๑", "")), widths=(80, 300))
    baseline = TOP + 38  # the second row's
    spilled = Boxed("ส่งรายงานความก้าวหน้าทุกหกเดือน", LEFT + 250, baseline - 1.1 * FONT_SIZE, baseline + 0.4 * FONT_SIZE, 0.99)

    pages = markdown_pages([grid], [grid, spilled])

    assert pages == ["[ตารางอาจไม่สมบูรณ์]\n\nงวดที่ งาน\n\n๑ ส่งรายงานความก้าวหน้าทุกหกเดือน"]


def test_a_box_reaching_past_the_rule_above_its_row_by_its_marks_keeps_the_table() -> None:
    # The engine boxes the last cell's line with its upper marks stacked high in a tight row: the box
    # is 2.2 font sizes tall and reaches 0.6 over the rule above it, as on a real scan, but holds no
    # line of the row above.
    grid = Grid(rows=(("งวดที่", "งาน"), ("๑", "")), widths=(80, 300))
    top = laid_out_grid(grid, TOP).cells[3].rect.y0
    tall = Boxed("ส่งรายงานทุกหกเดือน", LEFT + 90, top - 0.6 * FONT_SIZE, top + 1.6 * FONT_SIZE, 0.99)

    pages = markdown_pages([grid], [grid, tall])

    assert pages == ["| งวดที่ | งาน |\n| --- | --- |\n| ๑ | ส่งรายงานทุกหกเดือน |"]


def test_a_box_reaching_past_the_rule_beside_it_by_its_padding_keeps_the_table() -> None:
    # On a real scan in small print, a box reached 0.53 font sizes past the rule to the left of its
    # text; this one reaches 0.6.
    grid = Grid(rows=(("งวดที่", "งาน"), ("๑", "")), widths=(80, 300))
    cell = laid_out_grid(grid, TOP).cells[3].rect
    padded = Boxed("ส่งรายงานทุกหกเดือน", cell.x0 - 0.6 * FONT_SIZE, cell.y0 + 2, cell.y0 + 2 + 1.5 * FONT_SIZE, 0.99)

    pages = markdown_pages([grid], [grid, padded])

    assert pages == ["| งวดที่ | งาน |\n| --- | --- |\n| ๑ | ส่งรายงานทุกหกเดือน |"]


def test_text_within_a_table_that_no_cell_holds_makes_it_incomplete() -> None:
    # The scan lost the rules round the last cell, so no cell holds its text.
    grid = Grid(rows=(("งวดที่", "งาน"), ("๑", Cell("ส่งรายงานทุกหกเดือน", boxed=False))), widths=(80, 300))

    pages = markdown_pages([grid], [grid])

    assert pages == ["[ตารางอาจไม่สมบูรณ์]\n\nงวดที่ งาน\n\n๑ ส่งรายงานทุกหกเดือน"]


def test_rules_round_no_text_make_no_table() -> None:
    # Such as the printed dots of a watermark, taken for rules: no cell they make holds text.
    empty = Grid(rows=(("", "", ""), ("", "", "")), widths=(120, 120, 120))
    lines: list[str | Grid] = ["ผู้รับทุนต้องส่งรายงานทุกหกเดือน", empty, "ตามเงื่อนไขในสัญญา"]

    pages = markdown_pages(lines, lines)

    assert pages == ["ผู้รับทุนต้องส่งรายงานทุกหกเดือน\n\nตามเงื่อนไขในสัญญา"]


AFTER_THE_PICTURE = TOP + LINE_HEIGHT - FONT_SIZE + 200 + LINE_HEIGHT  # the baseline below a picture 200 pt tall


@pytest.mark.parametrize(
    "marking",
    [
        Picture("dark photo", 300, 200),
        Picture("stamp", 200, 200),
        Styled("ลงชื่อ ____________________________________ ผู้ให้ทุน"),
    ],
    ids=["a dark photo", "a stamp", "an underline"],
)
def test_a_marking_on_a_scan_that_is_not_a_table_makes_none(marking: Picture | Styled) -> None:
    after = Styled("ผู้รับทุนต้องส่งรายงานทุกหกเดือน", y=AFTER_THE_PICTURE)
    read = [Styled("หน้าที่มีภาพ"), *([] if isinstance(marking, Picture) else [marking]), after]

    [page] = markdown_pages([read[0], marking, after], read)

    assert "|" not in page and "<table>" not in page
    assert [line.text for line in read if line.text not in page] == []


def test_a_table_drawn_inside_a_cell_of_another_is_part_of_that_cell() -> None:
    # As on text pages, the inner grid is dropped and its text goes to the cell around it.
    outer = Grid(rows=(("หัวข้อ", "รายละเอียด"), ("แผนงาน", "\n\n\n")), widths=(120, 300))
    inner = Grid(rows=(("ปี", "งบ"), ("1", "100")), widths=(80, 80), x=220, y=TOP + 34)
    lines: list[str | Grid] = [outer, inner]

    pages = markdown_pages(lines, lines)

    assert pages == ["| หัวข้อ | รายละเอียด |\n| --- | --- |\n| แผนงาน | ปี งบ<br>1 100 |"]


def test_rules_that_leave_a_space_no_cell_could_be_make_an_incomplete_table() -> None:
    # Three cells are left unruled, so the rules around them close off an L: it cannot be told
    # which of them the text in it belongs to, so the rows are read as lines, after the label.
    def cell(row: int, column: int) -> Cell:
        return Cell(f"{row}{column}", boxed=(row, column) not in {(2, 2), (2, 3), (3, 2)})

    grid = Grid(rows=tuple(tuple(cell(row, column) for column in range(1, 5)) for row in range(1, 5)), widths=(90,) * 4)

    pages = markdown_pages([grid], [grid])

    assert pages == ["[ตารางอาจไม่สมบูรณ์]\n\n11 12 13 14\n\n21 22 23 24\n\n31 32 33 34\n\n41 42 43 44"]


def test_a_ruled_header_box_repeated_at_the_top_of_scans_is_page_furniture() -> None:
    header = Grid(rows=(("สัญญาเลขที่ T-001", "ฉบับลงนาม"), ("โครงการพัฒนาระบบ", "เอกสารลับ")), widths=(200, 150), y=36)
    pages: list[list[str | Styled | Grid]] = [
        [header, Styled("เนื้อหาหน้าแรกของสัญญา", y=140)],
        [header, Styled("เนื้อหาหน้าที่สองของสัญญา", y=140)],
    ]

    result = converted(pdf(*(scanned(page, seed=1) for page in pages)), FakeEngine(*pages))

    assert result == ["เนื้อหาหน้าแรกของสัญญา", "เนื้อหาหน้าที่สองของสัญญา"]


# Uncertain text in cells: marked and counted as it is in paragraphs


def test_a_word_misread_in_a_cell_is_marked() -> None:
    # The engine reads "เดือน" as "เคือน", which no word list holds.
    pages = markdown_pages([schedule("ส่งรายงานทุกหกเดือน")], [schedule("ส่งรายงานทุกหกเคือน")])

    assert pages == ["| งวดที่ | งาน |\n| --- | --- |\n| ๑ | ส่งรายงานทุกหก<mark>เคือน</mark> |"]


def test_neighbouring_misread_words_in_a_merged_table_are_marked_as_one() -> None:
    # "เดือน ศูนย์" read as two words that are not words, in a table with a merged cell.
    def merged(task: str) -> Grid:
        return Grid(rows=((Cell("งวดที่", rowspan=2), "งาน"), (task,), ("๑", "50,000")), widths=(80, 300))

    pages = markdown_pages([merged("ส่งรายงานทุกหกเดือน ศูนย์ตรวจรับแล้ว")], [merged("ส่งรายงานทุกหกเคือน ศนย์ตรวจรับแล้ว")])

    assert pages == [
        "<table>\n"
        '<tr><td rowspan="2">งวดที่</td><td>งาน</td></tr>\n'
        "<tr><td>ส่งรายงานทุกหก<mark>เคือน ศนย์</mark>ตรวจรับแล้ว</td></tr>\n"
        "<tr><td>๑</td><td>50,000</td></tr>\n"
        "</table>"
    ]


def test_a_mark_in_a_cell_keeps_its_bars_escaped_and_its_line_breaks() -> None:
    # Two short items, the first misread, beside a code with a bar in it.
    pages = markdown_pages(
        [Grid(rows=(("รหัส", "รายการ"), ("A|B", "รายงานรายเดือน\nเครื่องพิมพ์")), widths=(100, 300))],
        [Grid(rows=(("รหัส", "รายการ"), ("A|B", "รายงานรายเคือน\nเครื่องพิมพ์")), widths=(100, 300))],
    )

    assert pages == ["| รหัส | รายการ |\n| --- | --- |\n" + r"| A\|B | รายงานราย<mark>เคือน</mark><br>เครื่องพิมพ์ |"]


def test_cells_the_engine_reads_with_low_confidence_are_marked_whole() -> None:
    # As handwriting in a table's cells is read.
    grid = Grid(rows=(("ปีที่", "งบประมาณ"), ("2568", "50,000")), widths=(150, 150))

    pages = markdown_pages([grid], [Unsure(grid)])

    assert pages == [
        "| <mark>ปีที่</mark> | <mark>งบประมาณ</mark> |\n| --- | --- |\n| <mark>2568</mark> | <mark>50,000</mark> |"
    ]


def test_the_summary_counts_uncertain_text_in_cells() -> None:
    lines: list[Styled | Grid] = [schedule("ส่งรายงานทุกหกเดือน"), Styled("ผู้รับทุนต้องส่งรายงานตามกำหนด", y=TOP + 3 * LINE_HEIGHT)]
    read: list[Styled | Grid] = [schedule("ส่งรายงานทุกหกเคือน"), Styled("ผู้รับทุนต้องส่งรายงานตามกำหนค", y=TOP + 3 * LINE_HEIGHT)]

    result = convert(pdf(scanned(lines, seed=1)), ocr=FakeEngine(read))

    assert [page.uncertain for page in result.summary.pages] == [2]
    assert "ข้อความที่ไม่แน่ใจ 2 จุด ในหน้า 1" in result.output.decode("utf-8").split("<!-- page 1 -->")[0]


def test_uncertain_text_in_a_cell_is_highlighted_yellow_in_the_word_table() -> None:
    result = convert(pdf(scanned([schedule("ส่งรายงานทุกหกเดือน")], seed=1)), to="docx", ocr=FakeEngine([schedule("ส่งรายงานทุกหกเคือน")]))

    [[table]] = read_docx(result.output).pages
    assert isinstance(table, Table)
    shown = runs(table.cell(1, 1).paragraphs[0])
    assert "".join(run.text for run in shown) == "ส่งรายงานทุกหกเคือน"
    assert [run.text for run in shown if run.highlight] == ["เคือน"]
    assert {run.highlight for run in shown} == {None, "yellow"}
    # The highlighted word is set in the same Thai font as the rest.
    assert {(run.thai_font, run.thai_size, run.complex_script) for run in shown} == {("TH Sarabun New", 16, True)}


def test_a_note_typed_in_a_cell_is_never_marked() -> None:
    # A surname the word list does not hold, typed over the scan: the text layer holds it exactly.
    grid = Grid(rows=(("งวดที่", "ผู้ตรวจรับ"), ("๑", "")), widths=(80, 300))
    note = Styled("นายสมชาย ชุติมณฑน์", x=156, y=TOP + 31)
    misread_note = Styled("นายสมชาย ชุติมณฑบ์", x=156, y=TOP + 31)
    body = ["ผู้รับทุนตกลงดำเนินงานตามแผนที่กำหนดไว้ในสัญญาฉบับนี้"] * 6
    lines: list[str | Grid] = [grid, "", *body]

    pages = converted(pdf(scanned(lines, overlay=[note], seed=1)), FakeEngine([*lines, misread_note]))

    assert pages[0].startswith("| งวดที่ | ผู้ตรวจรับ |\n| --- | --- |\n| ๑ | นายสมชาย ชุติมณฑน์ |")


def test_each_item_kept_on_a_line_of_its_own_is_marked_apart() -> None:
    # Two short items read with low confidence: two stretches, not one across the line break.
    grid = Grid(rows=(("ปี", "รายการ"), ("2568", "เครื่องพิมพ์\nเครื่องสแกน")), widths=(100, 300))

    result = convert(pdf(scanned([grid], seed=1)), ocr=FakeEngine([Unsure(grid)]))

    assert "| <mark>2568</mark> | <mark>เครื่องพิมพ์</mark><br><mark>เครื่องสแกน</mark> |" in result.output.decode("utf-8")
    assert [page.uncertain for page in result.summary.pages] == [5]


# Boxed mark boxes in cells: a copy of a line's upper marks goes, but never across a rule


def test_a_copy_of_upper_marks_boxed_apart_in_a_cell_is_dropped() -> None:
    # The cell's text line sits 3 points into its cell, under the row's top rule: TOP + 31. The scan
    # is not askew, so that its rules lie where the engine's boxes say.
    lines: list[str | Grid] = [schedule("ส่งรายงานความก้าวหน้าทุกหกเดือน")]

    pages = converted(pdf(Scan(lines, dpi=200, noise=6, seed=1)), FakeEngine([*lines, marks_over(TOP + 31, x=156)]))

    assert pages == ["| งวดที่ | งาน |\n| --- | --- |\n| ๑ | ส่งรายงานความก้าวหน้าทุกหกเดือน |"]


def test_a_short_dash_alone_in_a_cell_stays() -> None:
    # Half way up the letters of the empty cell's line, whose baseline is TOP + 31.
    lines: list[str | Grid] = [schedule("")]

    pages = markdown_pages(lines, [*lines, Boxed("-", 156, TOP + 21, TOP + 27, confidence=0.99)])

    assert pages == ["| งวดที่ | งาน |\n| --- | --- |\n| ๑ | - |"]


def test_a_dash_at_the_foot_of_a_cell_is_never_taken_for_the_marks_of_the_line_in_the_cell_below() -> None:
    # The rows' rules lie at TOP + 12, + 40 and + 68. The line in the cell below, marked "ผู้", is
    # boxed from 4 points over its rule, as PaddleOCR boxes a marked line; the dash at the foot of
    # the cell above lies over the top of that box, over the marks, but across the rule.
    scanned_grid = Grid(rows=(("รายการ", "หมายเหตุ"), ("", ""), ("ค่าตอบแทนผู้วิจัย", "ตรวจแล้ว")), widths=(200, 200))
    read_grid = Grid(rows=(("รายการ", "หมายเหตุ"), ("", ""), ("", "ตรวจแล้ว")), widths=(200, 200))
    marked_line = Boxed("ค่าตอบแทนผู้วิจัย", LEFT - 1, TOP + 36, TOP + 65, confidence=0.99)
    dash = Boxed("-", 136, TOP + 32, TOP + 38, confidence=0.99)

    pages = converted(pdf(Scan([scanned_grid], dpi=200, noise=6, seed=1)), FakeEngine([read_grid, marked_line, dash]))

    assert pages == ["| รายการ | หมายเหตุ |\n| --- | --- |\n| - |  |\n| ค่าตอบแทนผู้วิจัย | ตรวจแล้ว |"]
